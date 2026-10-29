"""Fiche de collecte D2 RTE (demande de PTF, installation de consommation) et courriel d'envoi.

La D2 est un formulaire Word envoye par courriel au service raccordement consommateurs de RTE, avec la preuve
du virement de la somme forfaitaire et le schema unifilaire. Le modele (modeles/fiche_d2_modele.docx) est la D2
« mise au propre » des dossiers OM01, SLN01 et RSB01, videe de ses valeurs : memes cases, memes coches, plus
propre que la fiche vierge de RTE (case « Nouveau raccordement » manquante, trois cases au lieu de deux).

remplir() ecrit chaque case depuis la fiche du site et la demande ; une valeur requise absente est ecrite
« à compléter » surlignee en jaune, comme la date de signature. Rien ne part sans relecture : la fiche sort
en .docx, a relire, signer et convertir en PDF par le demandeur. courriel() prepare le message d'envoi
(.eml non envoye, X-Unsent) avec les pieces : le dernier clic, « Envoyer », reste humain.
"""

from __future__ import annotations

import copy
import io
import re
from datetime import date
from email.message import EmailMessage
from email.utils import formataddr, make_msgid
from pathlib import Path

MODELE = Path(__file__).resolve().parent / "modeles" / "fiche_d2_modele.docx"
DESTINATAIRE_RTE = "rte-fcent-raccordement-conso-dist@rte-france.com"
FORFAIT_KEUR = 42
A_COMPLETER = "à compléter"
TAILLE_MAX_COURRIEL = 15 * 1024 * 1024  # au-dela, la plupart des messageries refusent

# Ordre et libelles des pieces dans le courriel (types de pieces Amont).
PIECES_COURRIEL = [
    ("fiche_d2", "la fiche de collecte D2 « consommation » (article 1.2.5 de la DTR), reprenant l'ensemble des données de la demande"),
    ("unifilaire", "le schéma électrique unifilaire de principe"),
    ("plan_situation", "le plan de situation du site"),
    ("plan_implantation", "le plan d'implantation envisagée"),
    ("annexe_cartes", "l'annexe « Cartes, implantation et cadastre »"),
    ("extrait_cadastral", "l'extrait du plan cadastral des parcelles, avec le point de raccordement envisagé"),
    ("planning", "le planning prévisionnel du projet"),
    ("montee_en_charge", "le fichier de montée en charge progressive"),
    ("preuve_virement", "la preuve du virement de la somme forfaitaire"),
    ("mandat", "le mandat du consommateur"),
]
TYPES_JOINTS = [t for t, _ in PIECES_COURRIEL]

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ------------------------------------------------------------------ mise en forme des valeurs
def siren_lisible(s) -> str:
    c = re.sub(r"\D", "", str(s or ""))
    return " ".join(c[i:i + 3] for i in range(0, 9, 3)) if len(c) == 9 else str(s or "")


def date_fr(iso) -> str:
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(iso or ""))
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}" if m else str(iso or "")


def dms(lat: float, lon: float) -> str:
    def un(v: float, pos: str, neg: str) -> str:
        h = pos if v >= 0 else neg
        v = abs(v)
        d = int(v)
        m = int((v - d) * 60)
        s = (v - d - m / 60) * 3600
        return f"{d}°{m:02d}'{s:04.1f}\" {h}"
    return f"{un(lat, 'N', 'S')}, {un(lon, 'E', 'O')}"


def domaine(tension_kv) -> str:
    try:
        t = float(tension_kv)
    except (TypeError, ValueError):
        return ""
    if t >= 350:
        return f"HTB3 ({t:g} kV)"
    if t >= 130:
        return f"HTB2 ({t:g} kV)"
    if t >= 50:
        return f"HTB1 ({t:g} kV)"
    return f"HTA ({t:g} kV)"


def nb(v) -> str:
    if v in (None, ""):
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f"{int(f)}" if f.is_integer() else f"{f:g}".replace(".", ",")


def paliers(texte: str) -> list[tuple[str, float]]:
    """« 2029 : 80 ; 2030 : 160 ; 2031 : 240 » (MW cumules) -> [("2029", 80), ...]. Liste vide si illisible."""
    out = []
    for m in re.finditer(r"((?:fin |T\d )?\d{4})\s*[:=,-]\s*(\d+(?:[.,]\d+)?)\s*(?:MW)?", texte or ""):
        out.append((m.group(1), float(m.group(2).replace(",", "."))))
    return out


def lignes_montee_en_charge(texte: str) -> list[str]:
    p = paliers(texte)
    if not p:
        return [x.strip() for x in re.split(r"[\n;]+", texte or "") if x.strip()]
    lignes, avant = [], 0.0
    for i, (quand, cumul) in enumerate(p, 1):
        ajout = cumul - avant
        lignes.append(f"- palier {i} : {nb(ajout)} MW, {quand}" + (f" ({nb(cumul)} MW cumulés)" if i > 1 else ""))
        avant = cumul
    return lignes


# ------------------------------------------------------------------ donnees de la fiche
def donnees(v: dict, demande: dict, num_etude: str | None, pieces: list[dict], code: str) -> dict:
    """Toutes les cases de la D2. v : valeurs typees de la fiche du site ; demande : la demande RTE PTF."""
    p_mw = demande.get("puissance_mw") or v.get("puissance_mw")
    nb_alim = demande.get("alimentations_supplementaires")
    if nb_alim in (None, ""):
        nb_alim = v.get("alimentations_secours")
    try:
        nb_alim = int(float(nb_alim)) if nb_alim not in (None, "") else None
    except ValueError:
        nb_alim = None
    lat = demande.get("gps_lat") or v.get("lat")
    lon = demande.get("gps_lon") or v.get("lon")
    cp_ville = " ".join(x for x in (str(v.get("code_postal") or ""), str(v.get("commune") or "")) if x)
    if cp_ville and v.get("departement"):
        cp_ville += f" ({v['departement']})"
    societe = v.get("demandeur_societe") or ""
    rep = v.get("demandeur_representant") or ""
    secours = str(v.get("production_secours") or "").strip().lower()
    return {
        "num_etude": num_etude or "",
        "societe": societe,
        "forme_juridique": v.get("demandeur_forme_juridique") or "",
        "siren": siren_lisible(v.get("demandeur_siren")),
        "rcs": v.get("demandeur_rcs") or "",
        "representant": rep,
        "siege_rue": v.get("demandeur_siege_adresse") or "",
        "siege_cp_ville": v.get("demandeur_siege_cp_ville") or "",
        "siege_pays": v.get("demandeur_siege_pays") or ("France" if v.get("demandeur_siege_adresse") else ""),
        "telephone": v.get("interlocuteur_tel") or "",
        "courriel": v.get("interlocuteur_mail") or "",
        "site": v.get("nom_installation") or code,
        "siret_site": v.get("site_siret") or "",
        "adresse": v.get("adresse") or "",
        "cp_ville": cp_ville,
        "point_connexion": v.get("point_connexion") or "",
        "gps": (f"{float(lat):.5f}, {float(lon):.5f} / {dms(float(lat), float(lon))}" if lat not in (None, "") and lon not in (None, "") else ""),
        "parcelles": v.get("parcelles") or "",
        "surface_ha": v.get("surface_ha"),
        "date_mad": date_fr(v.get("date_mise_a_disposition")),
        "pracc": nb(p_mw),
        "psoutirage": nb(v.get("puissance_soutirage_mw") or p_mw),
        "pmax": nb(v.get("puissance_soutirage_mw") or p_mw),
        "secours": True if secours in ("oui", "o", "true", "1") else (False if secours in ("non", "n", "false", "0") else None),
        "secours_detail": v.get("secours_groupes") or "",
        "transfo_tensions": v.get("transfo_tensions") or "",
        "transfo_detail": v.get("transfo_detail") or "",
        "alim_supp": None if nb_alim is None else nb_alim > 0,
        "alim_1": nb(p_mw) if nb_alim and nb_alim >= 1 else "",
        "alim_2": nb(p_mw) if nb_alim and nb_alim >= 2 else "",
        "tension_kv": v.get("tension_kv"),
        "montee": v.get("montee_en_charge") or "",
        "reserve_rte": v.get("reserve_fonciere_rte") or "",
        "complements": v.get("complements_d2") or "",
        "pieces": [p["type"] for p in pieces],
        "signataire": (f"{rep} de {societe}" if rep and societe else rep),
    }


# Cases dont l'absence est signalee (ecrites « à compléter » en jaune dans la fiche).
REQUIS_D2 = {
    "societe": "société demanderesse", "forme_juridique": "forme juridique", "siren": "SIREN", "rcs": "lieu d'immatriculation RCS",
    "representant": "représentant légal", "siege_rue": "siège social, numéro et rue", "siege_cp_ville": "siège social, code postal et ville",
    "telephone": "téléphone", "courriel": "courriel", "adresse": "adresse du site", "cp_ville": "code postal et commune du site",
    "gps": "coordonnées GPS du point de connexion", "date_mad": "date souhaitée de mise à disposition", "pracc": "puissance de raccordement",
}


def manquants(x: dict) -> list[str]:
    out = [lib for k, lib in REQUIS_D2.items() if not x.get(k)]
    if x.get("secours") is None:
        out.append("production de secours, oui ou non")
    if x.get("alim_supp") is None:
        out.append("nombre d'alimentations HTB supplémentaires")
    return out


def elements_complementaires(x: dict) -> list[tuple[str, str]]:
    """Section VI : (style, texte). Styles : vide, corps, titre, puce."""
    out: list[tuple[str, str]] = [("vide", "")]
    montee = lignes_montee_en_charge(x["montee"])
    if montee:
        out += [("corps", "Montée en charge progressive :"), *[("corps", ligne) for ligne in montee], ("vide", "")]
    dom = domaine(x.get("tension_kv"))
    phrase = f"Dimensionnement du raccordement pour {x['pracc'] or '…'} MW" + (f" en {dom}" if dom else "")
    if x.get("alim_supp"):
        phrase += ", avec une alimentation supplémentaire de même puissance assurant la redondance (N+1)"
    out.append(("corps", phrase + "."))
    if montee:
        out.append(("corps", "Nous vous remercions d'étudier l'ensemble des solutions susceptibles de réduire le délai de "
                             "raccordement et de permettre une mise à disposition anticipée du premier palier de puissance, "
                             "dès que cela sera techniquement réalisable."))
    if x.get("reserve_rte"):
        out.append(("corps", "Afin d'optimiser ce calendrier, nous pouvons nous engager à mettre à la disposition de RTE le "
                             f"foncier nécessaire à l'implantation des ouvrages ({x['reserve_rte']})."))
    if x.get("complements"):
        out += [("corps", ligne.strip()) for ligne in x["complements"].splitlines() if ligne.strip()]
    libelles = dict(PIECES_COURRIEL)
    jointes = [t for t in TYPES_JOINTS if t in x["pieces"] and t not in ("fiche_d2", "preuve_virement", "mandat")]
    if jointes:
        out.append(("vide", ""))
        titre = "Pièces jointes"
        if x.get("num_etude"):
            titre += f", établies à la suite de l'étude exploratoire RTE n° {x['num_etude']}"
        out.append(("titre", titre + " :"))
        out += [("puce", f"• {libelles[t][0].upper()}{libelles[t][1:]} ;") for t in jointes]
        out[-1] = ("puce", out[-1][1][:-2] + ".")
    return out


# ------------------------------------------------------------------ ecriture du .docx
def _ecrire_paragraphe(p, texte: str, jaune: bool = False, taille: float | None = None, neutre: bool = False) -> None:
    """neutre : texte droit et noir, pour ecrire dans un paragraphe d'aide (gris italique) ou de pointilles."""
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor
    runs = p.runs
    if not runs:
        r = p.add_run(texte)
    else:
        r = runs[0]
        r.text = texte
        for x in runs[1:]:
            x._r.getparent().remove(x._r)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW if jaune else None
    if taille:
        r.font.size = Pt(taille)
    if neutre:
        r.font.italic = False
        r.font.color.rgb = RGBColor(0, 0, 0)


def _ecrire_case(cellule, valeur, requis: bool = False) -> None:
    """Remplace le contenu d'une case (premiere ligne gardee pour la mise en forme). Liste : une ligne par element."""
    import docx
    lignes = valeur if isinstance(valeur, list) else [str(valeur or "")]
    jaune = False
    if requis and not any(lignes):
        lignes, jaune = [A_COMPLETER], True
    ps = cellule.paragraphs
    modele = ps[0]
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    _ecrire_paragraphe(modele, lignes[0], jaune)
    dernier = modele
    for ligne in lignes[1:]:
        neuf = copy.deepcopy(modele._p)
        dernier._p.addnext(neuf)
        dernier = docx.text.paragraph.Paragraph(neuf, cellule)
        _ecrire_paragraphe(dernier, ligne, jaune)


# Cases OUI/NON. Les carres sont dessines (formes ancrees) ; le « X » est du texte tape dedans, comme dans les
# D2 envoyees (OM01, SLN01, RSB01). Non cochee : retrait de la fiche vierge RTE (libelle juste apres le carre).
# Cochee : sans retrait, espaces + « X » + espaces + libelle, cale a l'ecran sur chaque case (rendu verifie).
CASES = {  # position : (libelle, retrait non cochee, retrait cochee, texte avant le X, texte entre X et libelle)
    (0, 2, 8): ("OUI", {"right": "211"}, {"right": "211"}, "", "   "), (0, 2, 9): ("NON", {"right": "133"}, {"right": "133"}, "", "   "),
    (0, 3, 8): ("OUI", {"right": "211"}, {"right": "211"}, "", "   "), (0, 3, 9): ("NON", {"right": "133"}, {"right": "133"}, "", "   "),
    (1, 13, 10): ("OUI", {"left": "428"}, None, "     ", "  "), (1, 13, 14): ("NON", {"left": "426"}, None, "    ", "  "),
    (2, 8, 4): ("OUI", None, None, "    ", "   "), (2, 8, 6): ("NON", {"left": "585"}, None, "       ", "   "),
    (2, 12, 4): ("OUI", None, None, "    ", "   "), (2, 12, 6): ("NON", {"left": "487"}, None, "       ", "   "),
}


def _cocher(cellule, position: tuple, oui: bool) -> None:
    from docx.oxml.ns import qn
    libelle, retrait, retrait_coche, avant, apres = CASES[position]
    p = cellule._tc.find(qn("w:p"))
    runs = [r for r in p.findall(qn("w:r")) if r.find(qn("w:t")) is not None]
    gabarit = next((r for r in runs if (r.find(qn("w:t")).text or "").strip().endswith(libelle)), runs[-1])
    for r in runs:
        p.remove(r)
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = p.makeelement(qn("w:pPr"), {})
        p.insert(0, ppr)
    ind = ppr.find(qn("w:ind"))
    if ind is not None:
        ppr.remove(ind)
    attributs = (retrait_coche if oui else retrait) or {}
    if attributs:
        ind = ppr.makeelement(qn("w:ind"), {qn("w:" + k): v for k, v in attributs.items()})
        # ordre impose par le schema de w:pPr : w:ind avant w:jc, w:rPr, etc.
        suivants = {qn("w:" + n) for n in ("contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
                                           "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
                                           "sectPr", "pPrChange")}
        apres_ind = next((e for e in ppr if e.tag in suivants), None)
        if apres_ind is not None:
            apres_ind.addprevious(ind)
        else:
            ppr.append(ind)
    texte = f"{avant}X{apres}{libelle}" if oui else libelle
    r = copy.deepcopy(gabarit)
    t = r.find(qn("w:t"))
    t.text = texte
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    # les formes (carres) restent dans leur propre run, avant le texte
    p.append(r)


def _paire(doc, t: int, r: int, c_oui: int, c_non: int, valeur) -> None:
    case = doc.tables[t].rows[r].cells
    _cocher(case[c_oui], (t, r, c_oui), valeur is True)
    _cocher(case[c_non], (t, r, c_non), valeur is False)


def remplir(x: dict, modele: Path = MODELE) -> bytes:
    """La D2 remplie, en .docx."""
    import docx
    doc = docx.Document(str(modele))
    cell = lambda t, r, c: doc.tables[t].rows[r].cells[c]  # noqa: E731

    # I - type de demande
    p = cell(0, 2, 3).paragraphs[1]
    _ecrire_paragraphe(p, "Si OUI, indiquer le numéro de l’étude :  " + (x["num_etude"] or ""))
    _paire(doc, 0, 2, 8, 9, True if x["num_etude"] else None)
    _paire(doc, 0, 3, 8, 9, None)

    # II - demandeur
    for (r, c), cle in {(5, 4): "societe", (5, 7): "forme_juridique", (6, 4): "siren", (7, 4): "rcs", (8, 4): "representant",
                        (9, 4): "siege_rue", (10, 4): "siege_cp_ville", (11, 4): "siege_pays", (12, 4): "telephone",
                        (13, 4): "courriel"}.items():
        _ecrire_case(cell(0, r, c), x[cle], requis=cle in REQUIS_D2)
    # III - consommateur (vide : le demandeur agit pour son propre compte)
    for (r, c) in [(16, 4), (16, 7), (17, 4), (18, 4), (19, 4), (20, 4), (21, 4), (22, 4), (23, 4), (24, 4)]:
        _ecrire_case(cell(0, r, c), "")

    # IV - site
    _ecrire_case(cell(1, 1, 6), x["site"])
    _ecrire_case(cell(1, 2, 6), x["siret_site"])
    _ecrire_case(cell(1, 3, 6), x["adresse"], requis=True)
    _ecrire_case(cell(1, 4, 6), x["cp_ville"], requis=True)
    c = cell(1, 5, 6)
    aide = c.paragraphs[0]
    for q in c.paragraphs[1:]:
        q._p.getparent().remove(q._p)
    _ecrire_paragraphe(aide, "Joindre un extrait cadastral ou schéma à échelle adaptée des parcelles concernées et indiquer la position envisagée")
    point = x["point_connexion"] or ("Voir le plan cadastral joint. L’emplacement du point de connexion pourra être ajusté selon "
                                     "les besoins de RTE, afin de simplifier l’acheminement de la puissance et réduire le délai de raccordement.")
    emprise = ""
    if x["parcelles"]:
        emprise = "Emprise : parcelles " + x["parcelles"] + (f" (≈ {nb(x['surface_ha'])} ha)" if x.get("surface_ha") else "")
    dernier = aide
    for texte, jaune in [(point, False), (x["gps"] or A_COMPLETER + " : coordonnées GPS", not x["gps"]), (emprise, False)]:
        if not texte:
            continue
        neuf = copy.deepcopy(aide._p)
        dernier._p.addnext(neuf)
        dernier = docx.text.paragraph.Paragraph(neuf, c)
        _ecrire_paragraphe(dernier, texte, jaune, taille=10, neutre=True)
    _ecrire_case(cell(1, 6, 6), x["date_mad"], requis=True)

    # V - donnees techniques
    _ecrire_case(cell(1, 9, 13), x["pracc"], requis=True)
    _ecrire_case(cell(1, 10, 13), x["psoutirage"], requis=True)
    _ecrire_case(cell(1, 11, 13), x["pmax"], requis=True)
    _paire(doc, 1, 13, 10, 14, x["secours"])
    _ecrire_case(cell(1, 16, 5), x["secours_detail"])
    _ecrire_case(cell(2, 3, 1), x["transfo_detail"] or "Non défini à ce stade (postes de transformation non encore choisis)")
    _ecrire_case(cell(2, 4, 1), "Non défini à ce stade" if not x["transfo_detail"] else "Voir ci-dessus")
    _ecrire_case(cell(2, 5, 1), x["transfo_tensions"])
    _paire(doc, 2, 8, 4, 6, x["alim_supp"])
    _ecrire_case(cell(2, 10, 2), x["alim_1"])
    _ecrire_case(cell(2, 11, 2), x["alim_2"])
    _paire(doc, 2, 12, 4, 6, False)

    # VI - elements complementaires
    c = cell(3, 1, 0)
    ps = c.paragraphs
    entete, pointilles = ps[0], ps[1]
    modele_p = copy.deepcopy(pointilles._p)
    for q in ps[1:]:
        q._p.getparent().remove(q._p)
    dernier = entete
    tailles = {"vide": None, "corps": 11, "titre": 10, "puce": 10}
    for style, texte in elements_complementaires(x) + [("pointilles", None)]:
        neuf = copy.deepcopy(modele_p)
        dernier._p.addnext(neuf)
        dernier = docx.text.paragraph.Paragraph(neuf, c)
        if texte is not None:
            _ecrire_paragraphe(dernier, texte, taille=tailles.get(style), neutre=True)

    # signature : date a la main, en jaune
    _ecrire_case(cell(4, 1, 0), f"../../{date.today().year}")
    _ecrire_paragraphe(cell(4, 1, 0).paragraphs[0], f"../../{date.today().year}", jaune=True)
    _ecrire_case(cell(4, 1, 1), x["signataire"], requis=True)
    for el in list(cell(4, 1, 2)._tc.iter()):
        if el.tag in (W + "drawing", W + "pict", W + "object"):
            el.getparent().remove(el)

    sortie = io.BytesIO()
    doc.save(sortie)
    return sortie.getvalue()


# ------------------------------------------------------------------ courriel d'envoi
def courriel(v: dict, demande: dict, num_etude: str | None, pieces: list[dict], destinataires: list[str],
             copie: list[str] | None = None) -> dict:
    """Objet, corps et liste ordonnee des pieces a joindre. pieces : pieces Amont de la demande (derniere version par type)."""
    site = v.get("nom_installation") or demande.get("projet_code") or ""
    commune = v.get("commune") or demande.get("commune") or ""
    dep = v.get("departement") or demande.get("departement") or ""
    p_mw = nb(demande.get("puissance_mw") or v.get("puissance_mw"))
    dom = domaine(v.get("tension_kv"))
    nb_alim = demande.get("alimentations_supplementaires") or v.get("alimentations_secours")
    try:
        nb_alim = int(float(nb_alim or 0))
    except ValueError:
        nb_alim = 0
    lieu = f"{commune} ({dep})" if dep else commune
    objet = f"{lieu} – {site} : demande de Proposition Technique et Financière"
    choisies = choisir_pieces(pieces)
    libelles = dict(PIECES_COURRIEL)
    corps = ["Madame, Monsieur,", ""]
    phrase = (f"Je vous prie de bien vouloir trouver ci-joint notre demande de Proposition Technique et Financière pour le "
              f"projet de centre de données {site}, situé {('à ' + lieu) if lieu else ''}".rstrip() + ".")
    corps.append(phrase)
    corps.append("")
    demande_txt = f"La demande porte sur une puissance de raccordement de {p_mw or '…'} MW au soutirage"
    if dom:
        demande_txt += f" en {dom}"
    if nb_alim:
        demande_txt += (", avec une alimentation supplémentaire de même puissance assurant une redondance N+1" if nb_alim == 1
                        else f", avec {nb_alim} alimentations supplémentaires")
    if v.get("date_mise_a_disposition"):
        demande_txt += f", pour une mise à disposition souhaitée au {date_fr(v['date_mise_a_disposition'])}"
    demande_txt += "."
    if num_etude:
        demande_txt += f" Elle fait suite à l'étude exploratoire n° {num_etude}."
    corps += [demande_txt, ""]
    if choisies:
        corps.append("Le dossier comprend les pièces suivantes :")
        corps.append("")
        for i, p in enumerate(choisies, 1):
            corps.append(f"{i}. {libelles.get(p['type'], p['nom_fichier'])} ;")
        corps[-1] = corps[-1][:-2] + "."
        corps.append("")
    if demande.get("date_virement"):
        corps.append(f"La somme forfaitaire de {FORFAIT_KEUR} 000 € HT a été versée le {date_fr(demande['date_virement'])}.")
    elif any(p["type"] == "preuve_virement" for p in choisies):
        corps.append(f"La somme forfaitaire de {FORFAIT_KEUR} 000 € HT a été versée ; la preuve du virement est jointe.")
    else:
        corps.append(f"La somme forfaitaire de {FORFAIT_KEUR} 000 € HT sera réglée selon les modalités que vous voudrez bien nous indiquer.")
    corps.append("")
    if v.get("reserve_fonciere_rte"):
        corps += ["Nous vous confirmons par ailleurs notre disponibilité pour mettre à disposition de RTE le foncier nécessaire "
                  "à l'implantation des ouvrages, si cela peut contribuer à optimiser le calendrier de raccordement.", ""]
    corps += ["Nous restons à votre entière disposition pour tout complément d'information utile à l'instruction du dossier, "
              "et pouvons organiser un point technique à votre convenance.", "",
              "Vous en souhaitant bonne réception, je vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.", ""]
    nom = " ".join(x for x in (v.get("interlocuteur_prenom"), v.get("interlocuteur_nom")) if x)
    corps += [x for x in (nom, v.get("interlocuteur_fonction"), v.get("interlocuteur_tel"), v.get("interlocuteur_mail")) if x]
    taille = sum(p.get("taille_octets") or 0 for p in choisies)
    alertes = []
    if not any(p["type"] == "fiche_d2" for p in choisies):
        alertes.append("fiche D2 absente : la générer, la signer et la joindre en PDF")
    elif not any(p["type"] == "fiche_d2" and p["nom_fichier"].lower().endswith(".pdf") for p in choisies):
        alertes.append("la D2 jointe est le .docx non signé : joindre la version signée en PDF (type fiche_d2) avant l'envoi")
    if not any(p["type"] == "unifilaire" for p in choisies):
        alertes.append("schéma unifilaire absent")
    if not any(p["type"] == "preuve_virement" for p in choisies) and not demande.get("date_virement"):
        alertes.append(f"preuve du virement de {FORFAIT_KEUR} k€ HT absente : la D2 n'est recevable qu'avec le virement")
    if taille > TAILLE_MAX_COURRIEL:
        alertes.append(f"pièces de {taille // (1024 * 1024)} Mo au total : au-delà de 15 Mo, envoyer par lien de transfert")
    return {"a": destinataires, "cc": copie or [], "objet": objet, "corps": "\n".join(corps), "pieces": choisies,
            "taille_octets": taille, "alertes": alertes}


def choisir_pieces(pieces: list[dict]) -> list[dict]:
    """Derniere version de chaque type joignable ; pour la D2, la version PDF (signee) l'emporte sur le .docx."""
    dernieres: dict[str, dict] = {}
    for p in sorted(pieces, key=lambda x: (x.get("ajoute_le") or "", x["id"])):
        if p["type"] not in TYPES_JOINTS:
            continue
        if p["type"] == "fiche_d2" and "fiche_d2" in dernieres:
            deja_pdf = dernieres["fiche_d2"]["nom_fichier"].lower().endswith(".pdf")
            if deja_pdf and not p["nom_fichier"].lower().endswith(".pdf"):
                continue
        dernieres[p["type"]] = p
    return [dernieres[t] for t in TYPES_JOINTS if t in dernieres]


def eml(message: dict, expediteur: str, nom_expediteur: str, fichiers: list[tuple[str, bytes]]) -> bytes:
    """Brouillon .eml (X-Unsent : Outlook l'ouvre comme un nouveau message pret a envoyer)."""
    m = EmailMessage()
    m["X-Unsent"] = "1"
    if expediteur:
        m["From"] = formataddr((nom_expediteur, expediteur)) if nom_expediteur else expediteur
    m["To"] = ", ".join(message["a"])
    if message.get("cc"):
        m["Cc"] = ", ".join(message["cc"])
    m["Subject"] = message["objet"]
    m["Message-ID"] = make_msgid(domain=(expediteur.split("@")[-1] if "@" in (expediteur or "") else "amont.local"))
    m.set_content(message["corps"])
    import mimetypes
    for nom, contenu in fichiers:
        typ = mimetypes.guess_type(nom)[0] or "application/octet-stream"
        maj, mineur = typ.split("/", 1)
        m.add_attachment(contenu, maintype=maj, subtype=mineur, filename=nom)
    return bytes(m)
