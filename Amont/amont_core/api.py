"""API HTTP d'Amont (FastAPI) et service du front statique.

Configuration par variables d'environnement, toutes facultatives :
  AMONT_DB        chemin de la base SQLite        (defaut : donnees/amont.sqlite)
  AMONT_FICHIERS  dossier des pieces               (defaut : donnees/raccordement)
  AMONT_URL       adresse publique pour les liens  (defaut : http://localhost:8000)
  AMONT_ENSEIGNE  nom affiche dans la barre        (defaut : Castignac)
  AMONT_SOURCES   dossier des fichiers sources     (defaut : donnees/sources)
  ANTHROPIC_API_KEY, AMONT_MODELE, AMONT_IA : extraction par IA, voir amont_core/extraction.py
"""

from __future__ import annotations

import io
import os
import re
from pathlib import Path

from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import VERSION, alertes, cadastre, courriels, d1, d2, enedis, extraction, grdf, ia, redaction, regles, reponses
from . import ressources as ressources_defaut
from . import fiche as fiche_site
from .stockage import Registre

RACINE = Path(__file__).resolve().parent.parent
CHEMIN_DB = Path(os.environ.get("AMONT_DB", RACINE / "donnees" / "amont.sqlite"))
DOSSIER_FICHIERS = Path(os.environ.get("AMONT_FICHIERS", RACINE / "donnees" / "raccordement"))
URL_BASE = os.environ.get("AMONT_URL", "http://localhost:8000").rstrip("/")
ENSEIGNE = os.environ.get("AMONT_ENSEIGNE", "Castignac")
DOSSIER_SOURCES = Path(os.environ.get("AMONT_SOURCES", RACINE / "donnees" / "sources"))
TAILLE_MAX_COURRIEL = 30 * 1024 * 1024
EXTENSIONS_SOURCES = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".json", ".txt", ".md", ".csv", ".png", ".jpg", ".jpeg", ".webp"}

app = FastAPI(title="Amont", version=VERSION, docs_url="/api/docs", redoc_url=None)
registre = Registre(CHEMIN_DB)


def _mes_adresses() -> list[str]:
    """Adresses considerees comme « nous » pour deviner le sens d'un courriel depose."""
    adresses = list(regles.CONFIG.get("mes_adresses") or [])
    mail = fiche_site.demandeur_defaut().get("interlocuteur_mail")
    if mail and mail not in adresses:
        adresses.append(mail)
    return adresses


def auteur(req: Request) -> str:
    """Adresse posee par Cloudflare Access, sinon le nom choisi dans la barre du site (en-tete X-Amont-Auteur,
    encode en URI pour les accents), sinon « interne »."""
    from urllib.parse import unquote
    a = req.headers.get("Cf-Access-Authenticated-User-Email") or unquote(req.headers.get("X-Amont-Auteur") or "").strip()
    return a[:80] or "interne"


@app.exception_handler(ValueError)
async def _valeur(_, exc: ValueError):
    return JSONResponse(status_code=400, content={"erreur": str(exc)})


@app.exception_handler(KeyError)
async def _absent(_, exc: KeyError):
    return JSONResponse(status_code=404, content={"erreur": f"introuvable : {exc.args[0] if exc.args else ''}"})


# ------------------------------------------------------------------ referentiel
@app.get("/api/sante")
def sante():
    return {"application": "Amont", "version": VERSION, "enseigne": ENSEIGNE, **registre.compte(),
            "base": str(CHEMIN_DB), "fichiers": str(DOSSIER_FICHIERS)}


@app.get("/api/referentiel")
def referentiel():
    return {
        "enseigne": ENSEIGNE,
        "ordre_filieres": regles.ORDRE_FILIERES,
        "filieres_optionnelles": {f: f in regles.ORDRE_FILIERES for f in regles.FILIERES_OPTIONNELLES},
        "filieres": {k: {"libelle": v["libelle"], "court": v["court"], "dimension": v["dimension"],
                         "etapes": [{"cle": e[0], "libelle": e[1], "detail": e[2]} for e in v["etapes"]]}
                     for k, v in regles.FILIERES.items()},
        "hors_chaine": regles.HORS_CHAINE,
        "champs_filiere": regles.CHAMPS_FILIERE,
        "types_pieces": regles.TYPES_PIECES,
        "pieces_obligatoires": regles.PIECES_OBLIGATOIRES_POUR_DEPOT,
        "fiche": {"champs": [{"cle": c[0], "libelle": c[1], "groupe": c[2], "type": c[3], "aide": c[4]} for c in fiche_site.CHAMPS_FICHE],
                  "groupes": fiche_site.GROUPES, "requis": fiche_site.REQUIS_PAR_FILIERE, "canaux": fiche_site.CANAUX},
        "types_sources": extraction.TYPES_SOURCES,
        "sens_echange": regles.SENS_ECHANGE,
        "canaux_echange": regles.CANAUX_ECHANGE,
        "delai_relance_jours": regles.DELAI_RELANCE_JOURS,
        "mes_adresses": _mes_adresses(),
        "utilisateurs": regles.CONFIG.get("utilisateurs") or [],
        "types_ressource": ressources_defaut.TYPES_RESSOURCE,
        "url_d1": d1.URL_D1,
        "validation": {"requise": regles.VALIDATION_REQUISE, "valideurs": regles.VALIDEURS},
        "ia": ia.disponible(),
        "modele": ia.description() or None,
        "ia_vision": ia.vision_disponible(),
        "redaction": {k: v["limite"] for k, v in redaction.TEXTES.items()},
    }


@app.get("/api/ia/etat")
def etat_ia():
    """Diagnostic : IA configuree, serveur du modele joignable, modeles presents."""
    return ia.etat()


@app.post("/api/sites/{code}/rediger")
async def rediger(code: str, req: Request):
    """{"champ": "contexte" | "complements_techniques"} : redige le texte de la D1 depuis la fiche, par l'IA locale.
    Le texte entre dans la fiche a verifier (champ vide) ou en valeur divergente (champ deja rempli)."""
    corps = await req.json()
    champ = corps.get("champ") or ""
    if not registre.obtenir_projet(code):
        raise KeyError(code)
    import anyio
    valeurs = registre.valeurs_fiche(code)
    try:
        r = await anyio.to_thread.run_sync(lambda: redaction.rediger(champ, valeurs))
    except ia.ErreurIA as e:
        raise ValueError(f"IA indisponible : {e}") from None
    bilan = registre.importer_fiche(code, {champ: r["texte"]}, origine="rédaction IA", auteur=auteur(req))
    return {**r, **(bilan if isinstance(bilan, dict) else {})}


@app.get("/api/moi")
def moi(req: Request):
    return {"auteur": auteur(req)}


# --------------------------------------------------------------------- projets
@app.get("/api/projets")
def projets():
    out = []
    for p in registre.lister_projets():
        s = registre.synthese_projet(p["code"])
        s.pop("demandes", None)
        s["par_filiere"] = {f: [_resume(d) for d in ds] for f, ds in s["par_filiere"].items()}
        out.append(s)
    return out


@app.post("/api/projets", status_code=201)
async def creer_projet(req: Request):
    return registre.creer_projet(await req.json(), auteur(req))


@app.get("/api/projets/{code}")
def projet(code: str):
    s = registre.synthese_projet(code)
    s["journal"] = registre.journal(projet=code, limite=100)
    pieces = []
    for d in s["demandes"]:
        for p in registre.lister_pieces(d["id"]):
            p["filiere"] = d["filiere"]
            pieces.append(p)
    s["pieces"] = pieces
    ech = []
    for d in s["demandes"]:
        for e in d["echeances"]:
            ech.append({**e, "demande_id": d["id"], "filiere": d["filiere"]})
    ech.sort(key=lambda e: e["date"])
    s["echeancier"] = ech
    return s


@app.patch("/api/projets/{code}")
async def modifier_projet(code: str, req: Request):
    return registre.modifier_projet(code, await req.json(), auteur(req))


@app.put("/api/projets/{code}/filieres/{filiere}")
async def definir_filiere(code: str, filiere: str, req: Request):
    """{"etat": "sans_objet", "motif": "..."} ou {"etat": null} pour revenir a « a suivre »."""
    corps = await req.json()
    return registre.definir_filiere(code, filiere, corps.get("etat"), corps.get("motif") or "", auteur(req))


# -------------------------------------------------------------------- demandes
def _resume(d: dict) -> dict:
    cles = ("id", "projet_code", "filiere", "statut", "libelle_statut", "site", "commune", "departement",
            "date_depot", "prochaine_echeance", "alertes", "dimension", "domaine", "puissance_mw",
            "hypothese_secondaire_mw", "puissance_kva", "debit_nm3h", "type_enedis", "num_etude",
            "ref_dossier", "cout_client_keur", "modifie_le", "modifie_par", "action", "echanges_resume",
            "prochaine_action")
    return {k: d.get(k) for k in cles if k in d}


@app.get("/api/demandes")
def demandes(filiere: str | None = None, projet: str | None = None, statut: str | None = None):
    return [_resume(d) for d in registre.lister_demandes(filiere, projet, statut)]


@app.post("/api/demandes", status_code=201)
async def creer_demande(req: Request):
    return registre.creer_demande(await req.json(), auteur(req))


@app.get("/api/demandes/{id_}")
def demande(id_: str):
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    d["pieces_manquantes"] = registre.pieces_manquantes(id_, d["filiere"])
    return d


@app.patch("/api/demandes/{id_}")
async def modifier_demande(id_: str, req: Request):
    corps = await req.json()
    origine = "robot" if req.headers.get("X-Amont-Origine") == "robot" else "dashboard"
    if origine == "robot" and corps.get("statut") == "depose":
        raise ValueError("un robot ne peut pas passer une demande à « déposé », c'est le geste humain")
    return registre.modifier_demande(id_, corps, auteur(req), origine)


@app.delete("/api/demandes/{id_}", status_code=204)
def supprimer_demande(id_: str, req: Request):
    if not registre.obtenir_demande(id_):
        raise KeyError(id_)
    registre.supprimer_demande(id_, auteur(req))


@app.get("/api/demandes/{id_}/payload")
def payload(id_: str):
    return registre.payload(id_)


@app.get("/api/demandes/{id_}/d1")
def plan_d1(id_: str):
    """Plan de remplissage du formulaire D1 du portail RTE, lu par le robot du poste."""
    p = registre.payload(id_)
    if p["filiere"] != "rte_ee":
        raise ValueError("le formulaire D1 ne concerne que les demandes RTE étude exploratoire")
    return d1.plan(p)


@app.get("/api/demandes/{id_}/enedis")
def plan_enedis(id_: str):
    """Plan de saisie du Portail Raccordement Enedis (parcours HTA), lu par le site et par le robot Enedis du poste."""
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    if d["filiere"] != "enedis":
        raise ValueError("plan Enedis : demandes Enedis seulement")
    return enedis.plan(registre.payload(id_))


@app.get("/api/demandes/{id_}/journal")
def journal_demande(id_: str):
    return registre.journal(demande_id=id_, limite=500)



# ----------------------------------------------------------- sites et sources
@app.get("/api/sites")
def sites():
    return registre.resume_sites()


@app.get("/api/sites/{code}")
def site(code: str):
    s = registre.obtenir_site(code)
    s["extrait_cadastral"] = _dernier_extrait(code)
    s["plan_situation"] = _dernier_extrait(code, "plan_situation")
    return s


def _dernier_extrait(code: str, motif: str = "extrait_cadastral") -> dict | None:
    """Dernier document cadastral (extrait ou plan de situation) recupere pour le site par Amont."""
    for x in registre.lister_sources(code):
        if motif in x["nom_fichier"] and x.get("type_source") == "cadastre":
            return {"source_id": x["id"], "nom_fichier": x["nom_fichier"], "le": x["ajoute_le"], "par": x["ajoute_par"],
                    "taille_octets": x["taille_octets"]}
    return None


def _ranger_dans_sources(code: str, nom_fichier: str, contenu: bytes, qui: str, type_source: str, methode: str) -> dict:
    dossier = DOSSIER_SOURCES / code
    dossier.mkdir(parents=True, exist_ok=True)
    cible = dossier / nom_fichier
    i = 1
    while cible.exists():
        cible = dossier / f"{Path(nom_fichier).stem}_{i}{Path(nom_fichier).suffix}"
        i += 1
    cible.write_bytes(contenu)
    src = registre.ajouter_source(code, cible.name, str(cible.relative_to(DOSSIER_SOURCES)).replace("\\", "/"), len(contenu), qui)
    registre.marquer_source(src["id"], type_source=type_source, statut="fait", methode=methode, nb_propositions=0,
                            extrait_le=src["ajoute_le"])
    return registre.obtenir_source(src["id"])


def _joindre(demande_id: str, type_piece: str, nom_fichier: str, contenu: bytes, qui: str) -> dict:
    d = registre.obtenir_demande(demande_id)
    dos = DOSSIER_FICHIERS / d["projet_code"] / d["filiere"]
    dos.mkdir(parents=True, exist_ok=True)
    fp = dos / nom_fichier
    j = 1
    while fp.exists():
        fp = dos / f"{Path(nom_fichier).stem}_{j}{Path(nom_fichier).suffix}"
        j += 1
    fp.write_bytes(contenu)
    return registre.ajouter_piece(demande_id, type_piece, fp.name, str(fp.relative_to(DOSSIER_FICHIERS)).replace("\\", "/"),
                                  len(contenu), qui)


def document_cadastre(code: str, nature: str, corps: dict, qui: str) -> dict:
    """Recupere sur cadastre.gouv.fr l'extrait (nature « extrait ») ou le plan de situation (« situation ») du site,
    le range dans les sources et le joint aux demandes RTE etude exploratoire. Bloquant : 20 s a 1 min."""
    if not registre.obtenir_projet(code):
        raise KeyError(code)
    v = registre.valeurs_fiche(code)
    args = (str(v.get("parcelles") or ""), str(v.get("commune") or ""), str(v.get("departement") or ""),
            str(v.get("code_insee") or ""), str(v.get("communes_secondaires") or ""))
    try:
        if nature == "situation":
            r = cadastre.plan_situation_site(*args, nom_site=str(v.get("nom_installation") or code), journal=lambda *a: None)
        elif nature == "gaz":
            # plan cadastral demande par GRDF : portrait, 1/1000e, impression « outils avancés » de cadastre.gouv.fr
            r = cadastre.extrait_site(*args, centrer_sur=corps.get("centrer_sur") or None,
                                      echelle=int(corps["echelle"]) if corps.get("echelle") else 1000,
                                      format_=corps.get("format") or "A4", orientation="Portrait",
                                      nom_site=str(v.get("nom_installation") or code), journal=lambda *a: None,
                                      prefixe_nom="plan_cadastral_1000")
        else:
            r = cadastre.extrait_site(*args, centrer_sur=corps.get("centrer_sur") or None,
                                      echelle=int(corps["echelle"]) if corps.get("echelle") else None,
                                      format_=corps.get("format") or None, nom_site=str(v.get("nom_installation") or code),
                                      journal=lambda *a: None)
    except cadastre.ErreurCadastre as e:
        raise ValueError(str(e)) from None
    src = _ranger_dans_sources(code, r.nom_fichier, r.pdf, qui, "cadastre", "cadastre.gouv.fr")
    type_piece = {"situation": "plan_situation", "gaz": "plan_cadastral_1000"}.get(nature, "extrait_cadastral")
    pieces, assemblages = [], []
    if corps.get("joindre", True):
        if nature == "gaz":
            cibles = [d for d in registre.lister_demandes(projet=code, filiere="grdf") if d["statut"] in regles.ETAPES_AVANT_ENVOI]
        else:
            cibles = registre.lister_demandes(projet=code, filiere="rte_ee") + [
                d for d in registre.lister_demandes(projet=code, filiere="enedis") if d["statut"] in regles.ETAPES_AVANT_ENVOI]
        for d in cibles:
            pieces.append(_joindre(d["id"], type_piece, Path(src["chemin"]).name, r.pdf, qui))
            if nature == "situation" and d["filiere"] == "rte_ee":
                try:
                    assemblages.append(assembler_plan_installation(d["id"], qui))
                except ValueError:
                    pass
    # ce qui a ete devine est signale a la verification (doublons de reference, parcelles introuvables)
    a_verifier = [a for a in r.avertissements if "existe à" in a or "commune principale" in a or "débordent" in a]
    if r.absentes:
        a_verifier.append("parcelles introuvables sur cadastre.gouv.fr : " + ", ".join(r.absentes))
    point = None
    if a_verifier and nature != "situation":
        point = registre.ajouter_commentaire({"projet_code": code, "cible": "fiche:parcelles", "a_verifier": True,
                                              "texte": "Extrait cadastral : " + " ; ".join(a_verifier)}, qui)
    return {"bilan": r.bilan(), "avertissements": r.avertissements, "absentes": r.absentes, "pages": r.pages,
            "communes": r.communes, "source": src, "pieces": pieces, "plans_installation": assemblages,
            "point_a_verifier": point["id"] if point else None, "taille_octets": len(r.pdf)}


ORDRE_PLAN_INSTALLATION = ["plan_situation", "plan_implantation", "unifilaire"]


def assembler_plan_installation(demande_id: str, qui: str) -> dict:
    """Le portail RTE n'a qu'un emplacement « plan de l'installation » (un PDF, 2 Mo) : on y fusionne la derniere
    version de chaque piece disponible, dans l'ordre plan de situation, plan d'implantation, schema unifilaire."""
    d = registre.obtenir_demande(demande_id)
    if not d:
        raise KeyError(demande_id)
    dernieres = {}
    for p in sorted(registre.lister_pieces(demande_id), key=lambda x: (x["ajoute_le"] or "", x["id"])):
        if p["type"] in ORDRE_PLAN_INSTALLATION:
            dernieres[p["type"]] = p
    choisies = [dernieres[t] for t in ORDRE_PLAN_INSTALLATION if t in dernieres]
    if not choisies:
        raise ValueError("aucune pièce à assembler : plan de situation, plan d'implantation ou schéma unifilaire")
    pdfs = []
    for p in choisies:
        chemin = DOSSIER_FICHIERS / p["chemin"]
        if chemin.suffix.lower() != ".pdf":
            raise ValueError(f"{p['nom_fichier']} n'est pas un PDF")
        pdfs.append(chemin.read_bytes())
    contenu = pdfs[0] if len(pdfs) == 1 else cadastre.assembler(pdfs)
    if len(contenu) > cadastre.TAILLE_MAX_RTE:
        raise ValueError(f"plan de l'installation de {len(contenu) // 1024} ko, au-delà des 2 Mo du portail RTE : alléger les pièces")
    from datetime import date as _date
    nom = f"{_date.today().strftime('%Y_%m_%d')}_plan_installation_{d['projet_code']}.pdf"
    piece = _joindre(demande_id, "plan_installation", nom, contenu, qui)
    return {"demande_id": demande_id, "piece": piece, "contenu": [p["type"] for p in choisies], "taille_octets": len(contenu)}


async def _corps(req: Request) -> dict:
    try:
        c = await req.json()
        return c if isinstance(c, dict) else {}
    except ValueError:
        return {}


@app.post("/api/sites/{code}/extrait-cadastral")
async def extrait_cadastral(code: str, req: Request):
    """Extrait du plan cadastral du site (cadastre.gouv.fr), d'apres les parcelles de la fiche.
    Corps facultatif : {"centrer_sur": "AH 367", "echelle": 2500, "format": "A3", "joindre": true}."""
    import anyio
    corps, qui = await _corps(req), auteur(req)
    return await anyio.to_thread.run_sync(lambda: document_cadastre(code, "extrait", corps, qui))


@app.put("/api/filieres/{filiere}")
async def activer_filiere(filiere: str, req: Request):
    """{"actif": true} : active ou retire une filiere optionnelle (NaTran) pour tout le site web."""
    corps = await _corps(req)
    return {"ordre_filieres": registre.definir_filiere_optionnelle(filiere, bool(corps.get("actif")), auteur(req))}


@app.post("/api/sites/{code}/ouvrir-natran", status_code=201)
def ouvrir_natran(code: str, req: Request):
    """GRDF redirige vers NaTran : active la filiere au besoin et cree la demande NaTran du site, pre-remplie depuis la fiche."""
    if not registre.obtenir_projet(code):
        raise KeyError(code)
    qui = auteur(req)
    if "natran" not in regles.ORDRE_FILIERES:
        registre.definir_filiere_optionnelle("natran", True, qui)
    cree = registre.preparer_demandes(code, ["natran"], qui)
    existantes = registre.lister_demandes(projet=code, filiere="natran")
    return {"cree": bool(cree), "demande": (cree or existantes)[0]}


@app.post("/api/sites/{code}/plan-cadastral-gaz")
async def plan_cadastral_gaz(code: str, req: Request):
    """Plan cadastral au format demande par la fiche GRDF : portrait, 1/1000e (A4 par defaut), joint aux demandes GRDF.
    Corps facultatif : {"centrer_sur": "AH 367", "format": "A3", "echelle": 1000, "joindre": true}. Un site plus grand que
    le cadre (200 m x 210 m en A4) se centre sur la parcelle du poste gaz."""
    import anyio
    corps, qui = await _corps(req), auteur(req)
    return await anyio.to_thread.run_sync(lambda: document_cadastre(code, "gaz", corps, qui))


@app.post("/api/sites/{code}/plan-situation")
async def plan_situation(code: str, req: Request):
    """Plan de situation DGFiP du site (toutes communes sur une page), joint aux demandes RTE EE, puis
    plan de l'installation reassemble."""
    import anyio
    corps, qui = await _corps(req), auteur(req)
    return await anyio.to_thread.run_sync(lambda: document_cadastre(code, "situation", corps, qui))


@app.post("/api/demandes/{id_}/validation/demander")
async def demander_validation(id_: str, req: Request):
    """Fige l'empreinte du dossier et cree un point « a verifier » pour le responsable. {"message": "..."}"""
    corps = await _corps(req)
    return registre.demander_validation(id_, corps.get("message") or "", auteur(req))


@app.post("/api/demandes/{id_}/validation/{decision}")
async def traiter_validation(id_: str, decision: str, req: Request):
    """decision : valider ou refuser. Reserve aux valideurs de amont_config.json. {"commentaire": "..."}"""
    corps = await _corps(req)
    dec = {"valider": "validee", "refuser": "refusee"}.get(decision)
    if not dec:
        raise ValueError("décision : valider ou refuser")
    return registre.traiter_validation(id_, dec, corps.get("commentaire") or "", auteur(req))


# ---------------------------------------------------------------- tout preparer
TACHES: dict[str, dict] = {}  # en memoire : suivi d'une preparation en cours, perdu au redemarrage (sans importance)


def _preparer_site(tache: dict, code: str, qui: str) -> None:
    """Enchaine les etapes de preparation d'un site. Une etape en echec n'arrete pas les suivantes."""
    import time as _t

    def etape(nom: str, fonction) -> None:
        e = {"nom": nom, "statut": "en_cours", "detail": ""}
        tache["etapes"].append(e)
        try:
            r = fonction()
            e["statut"], e["detail"] = (r if isinstance(r, tuple) else ("fait", r or ""))
        except Exception as ex:  # noqa: BLE001 : on note et on continue
            e["statut"], e["detail"] = "erreur", str(ex)

    v = registre.valeurs_fiche(code)
    so = registre.filieres_etat(code) if hasattr(registre, "filieres_etat") else {}

    def demande_ee():
        if "rte_ee" in so:
            return "ignore", "RTE étude exploratoire marquée sans objet pour ce site"
        ds = registre.lister_demandes(projet=code, filiere="rte_ee")
        if ds:
            return "deja", f"{ds[0]['id']} existe ({ds[0]['libelle_statut'].lower()})"
        cree = registre.preparer_demandes(code, ["rte_ee"], qui)
        return "fait", f"{cree[0]['id']} créée, pré-remplie depuis la fiche"
    etape("Demande RTE étude exploratoire", demande_ee)
    ee = registre.lister_demandes(projet=code, filiere="rte_ee")

    def sources():
        src = registre.lister_sources(code)
        en_cours = [x for x in src if x["statut"] in ("en_attente", "en_cours")]
        if en_cours:
            debut = _t.time()
            while en_cours and _t.time() - debut < 900:
                _t.sleep(3)
                en_cours = [x for x in registre.lister_sources(code) if x["statut"] in ("en_attente", "en_cours")]
            if en_cours:
                return "erreur", f"{len(en_cours)} fichier(s) encore en lecture après 15 min"
            return "fait", "lecture des fichiers terminée"
        lus = [x for x in src if x["statut"] == "fait" and x.get("type_source") != "cadastre"]
        return ("deja" if lus else "ignore"), (f"{len(lus)} fichier(s) déjà lus" if lus else "aucun fichier déposé dans les sources")
    etape("Lecture des fichiers du site", sources)

    def avec_pieces(type_piece: str, nature: str):
        if not ee:
            return "ignore", "pas de demande RTE EE"
        if not (registre.valeurs_fiche(code).get("parcelles")):
            return "erreur", "parcelles absentes de la fiche du site"
        if all(any(p["type"] == type_piece for p in registre.lister_pieces(d["id"])) for d in ee):
            return "deja", "déjà joint à " + ", ".join(d["id"] for d in ee)
        r = document_cadastre(code, nature, {}, qui)
        return "fait", r["bilan"] + ("" if not r["avertissements"] else " ⚠ " + " ; ".join(r["avertissements"]))
    etape("Extrait du plan cadastral", lambda: avec_pieces("extrait_cadastral", "extrait"))
    etape("Plan de situation et plan de l'installation", lambda: avec_pieces("plan_situation", "situation"))

    def textes():
        if not ia.disponible():
            return "ignore", "IA non configurée"
        faits = []
        for champ in redaction.TEXTES:
            if registre.valeurs_fiche(code).get(champ):
                continue
            r = redaction.rediger(champ, registre.valeurs_fiche(code))
            registre.importer_fiche(code, {champ: r["texte"]}, origine="rédaction IA", auteur=qui)
            faits.append(f"{fiche_site.LIBELLES.get(champ, champ)} ({r['longueur']} car.)"
                         + (f", nombres à vérifier : {', '.join(r['nombres_a_verifier'])}" if r["nombres_a_verifier"] else ""))
        return ("fait", "rédigé : " + " ; ".join(faits)) if faits else ("deja", "textes déjà présents dans la fiche")
    etape("Textes de la D1 (IA)", textes)

    def fiche_d2():
        ptf = registre.lister_demandes(projet=code, filiere="rte_ptf")
        if not ptf:
            return "ignore", "pas de demande RTE PTF pour ce site"
        faits = []
        for d in ptf:
            if any(p["type"] == "fiche_d2" for p in registre.lister_pieces(d["id"])):
                faits.append(f"{d['id']} : D2 déjà jointe")
                continue
            if d["statut"] not in regles.ETAPES_AVANT_ENVOI:
                faits.append(f"{d['id']} : déjà partie")
                continue
            from datetime import date as _date
            x = _donnees_d2(d)
            nom = f"{_date.today().strftime('%Y_%m_%d')}_fiche_D2_{code}.docx"
            _joindre(d["id"], "fiche_d2", nom, d2.remplir(x), qui)
            manque = d2.manquants(x)
            faits.append(f"{d['id']} : D2 générée" + (f", à compléter : {', '.join(manque)}" if manque else ", complète"))
        return ("fait" if any("générée" in f for f in faits) else "deja"), " ; ".join(faits)
    etape("Fiche D2 (RTE PTF)", fiche_d2)

    def fiche_grdf():
        if "grdf" in so:
            return "ignore", "GRDF marquée sans objet pour ce site"
        gaz = registre.lister_demandes(projet=code, filiere="grdf")
        faits = []
        if not gaz:
            cree = registre.preparer_demandes(code, ["grdf"], qui)
            gaz = cree
            faits.append(f"{cree[0]['id']} créée, pré-remplie depuis la fiche")
        from datetime import date as _date
        for d in gaz:
            if any(p["type"] == "fiche_grdf" for p in registre.lister_pieces(d["id"])):
                faits.append(f"{d['id']} : fiche GRDF déjà jointe")
                continue
            if d["statut"] not in regles.ETAPES_AVANT_ENVOI:
                faits.append(f"{d['id']} : déjà partie")
                continue
            x = _donnees_grdf(d)
            nom = f"{_date.today().strftime('%Y_%m_%d')}_fiche_GRDF_{code}.pdf"
            _joindre(d["id"], "fiche_grdf", nom, grdf.remplir(x), qui)
            manque = grdf.manquants(x)
            faits.append(f"{d['id']} : fiche GRDF générée" + (f", à compléter : {', '.join(manque)}" if manque else ", complète"))
        return ("fait" if any("générée" in f or "créée" in f for f in faits) else "deja"), " ; ".join(faits)
    etape("Fiche GRDF Grands Projets", fiche_grdf)

    def plan_gaz():
        gaz = [d for d in registre.lister_demandes(projet=code, filiere="grdf") if d["statut"] in regles.ETAPES_AVANT_ENVOI]
        if not gaz:
            return "ignore", "pas de demande GRDF à préparer"
        if not (registre.valeurs_fiche(code).get("parcelles")):
            return "erreur", "parcelles absentes de la fiche du site"
        if all(any(p["type"] == "plan_cadastral_1000" for p in registre.lister_pieces(d["id"])) for d in gaz):
            return "deja", "plan cadastral 1/1000e déjà joint"
        r = document_cadastre(code, "gaz", {}, qui)
        return "fait", r["bilan"] + ("" if not r["avertissements"] else " ⚠ " + " ; ".join(r["avertissements"]))
    etape("Plan cadastral 1/1000e pour GRDF", plan_gaz)

    def bilan():
        lignes = []
        site = registre.obtenir_site(code)
        av = sum(1 for x in site["fiche"].values() if x.get("valeur") not in (None, "") and not x.get("valide"))
        lignes.append(f"fiche : {av} valeur(s) à vérifier, {len(site['propositions'])} divergence(s)")
        for d in registre.lister_demandes(projet=code, filiere="rte_ee"):
            pl = d1.plan(registre.payload(d["id"]))
            lignes.append(f"{d['id']} : " + ("formulaire D1 complet" if pl["pret"] else "manque " + " ; ".join(pl["manquants"]))
                          + f" ; validation {d['validation']['etat']} ; prochaine action : {d['action']['texte']}")
        return "fait", " | ".join(lignes)
    etape("Bilan", bilan)
    tache["fini"] = True


@app.post("/api/sites/{code}/tout-preparer")
def tout_preparer(code: str, req: Request):
    """Lance la preparation du site en tache de fond ; suivre avec GET /api/taches/{id}."""
    import threading
    import uuid
    if not registre.obtenir_projet(code):
        raise KeyError(code)
    for t in TACHES.values():
        if t["code"] == code and not t["fini"]:
            return t
    tache = {"id": uuid.uuid4().hex[:10], "code": code, "etapes": [], "fini": False, "debut": __import__("datetime").datetime.now().isoformat(timespec="seconds")}
    TACHES[tache["id"]] = tache
    threading.Thread(target=_preparer_site, args=(tache, code, auteur(req)), daemon=True).start()
    return tache


@app.get("/api/taches/{tid}")
def tache(tid: str):
    if tid not in TACHES:
        raise KeyError(tid)
    return TACHES[tid]


# ------------------------------------------------- reponses des operateurs
def _texte_pdf(chemin: Path) -> str:
    """Texte avec mise en page conservee (les regles de lecture des etudes RTE s'appuient dessus)."""
    import shutil
    import subprocess
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(chemin), "-"], capture_output=True, timeout=120)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.decode("utf-8", "replace")
    return extraction.lire_texte(chemin)


def lire_piece(piece_id: int, qui: str) -> dict:
    r = registre.cx.execute("SELECT * FROM pieces WHERE id = ?", (piece_id,)).fetchone()
    if not r:
        raise KeyError(piece_id)
    chemin = DOSSIER_FICHIERS / r["chemin"]
    if chemin.suffix.lower() != ".pdf":
        raise ValueError("lecture des réponses : PDF seulement")
    texte = _texte_pdf(chemin)
    filiere = registre.obtenir_demande(r["demande_id"])["filiere"]
    if reponses.est_ptf_rte(texte) and filiere != "rte_ptf":
        raise ValueError("ce document est une PTF RTE : le déposer sur la demande RTE PTF du projet")
    if filiere == "rte_ptf" and not reponses.est_ptf_rte(texte):
        raise ValueError("ce document n'est pas reconnu comme une PTF RTE (conditions particulières)")
    if reponses.est_ptf_rte(texte):
        analyse = reponses.lire_ptf_complete(texte)
        if not analyse.get("num_ptf") and analyse.get("delai_mois") is None:
            raise ValueError("PTF reconnue mais illisible : gabarit inhabituel, à reporter à la main")
        return registre.ajouter_analyse(r["demande_id"], piece_id, analyse["type"], analyse, qui)
    if not reponses.est_etude_rte(texte):
        raise ValueError("ce document n'est reconnu ni comme une étude exploratoire RTE, ni comme une PTF RTE")
    analyse = reponses.lire(texte)
    if not analyse["solutions"]:
        raise ValueError("étude reconnue mais aucune solution lue : gabarit inhabituel, à reporter à la main")
    return registre.ajouter_analyse(r["demande_id"], piece_id, analyse["type"], analyse, qui)


@app.post("/api/pieces/{piece_id}/lire")
def lire_reponse(piece_id: int, req: Request):
    """Lit une piece deja jointe (etude exploratoire RTE) et propose ses solutions."""
    return lire_piece(piece_id, auteur(req))


@app.post("/api/demandes/{id_}/reponse", status_code=201)
async def deposer_reponse(id_: str, req: Request, fichier: UploadFile = File(...)):
    """Depot de l'etude recue : rangee en piece « etude_recue », puis lue."""
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    contenu = await fichier.read()
    if len(contenu) > 25 * 1024 * 1024:
        raise ValueError("fichier au-delà de 25 Mo")
    nom = re.sub(r"[^A-Za-z0-9_.\-]", "_", Path(fichier.filename or "etude.pdf").name)
    piece = _joindre(id_, "ptf" if d["filiere"] == "rte_ptf" else "etude_recue", nom, contenu, auteur(req))
    try:
        return {"piece": piece, "analyse": lire_piece(piece["id"], auteur(req))}
    except ValueError as e:
        return {"piece": piece, "analyse": None, "erreur": str(e)}


@app.post("/api/analyses/{aid}/reporter")
async def reporter_analyse(aid: int, req: Request):
    """{"solution": 1, "passer_etude_recue": true} : reporte numero, date, distance et la solution choisie dans la demande,
    passe la demande a « etude recue » si elle attendait l'etude, et cree un point a verifier."""
    corps = await _corps(req)
    a = registre.obtenir_analyse(aid)
    if not a:
        raise KeyError(aid)
    qui = auteur(req)
    if a["type"] == "ptf_rte":
        return _reporter_ptf(a, qui)
    num = corps.get("solution")
    num = int(num) if num not in (None, "") else None
    champs = reponses.champs_demande(a["donnees"], num)
    d = registre.obtenir_demande(a["demande_id"])
    statut = "etude_recue" if corps.get("passer_etude_recue", True) and d["statut"] in ("depose", "recevable", "pret_validation") else None
    if not d.get("date_depot") and a["donnees"].get("date_demande"):
        champs["date_depot"] = a["donnees"]["date_demande"]
    d = registre.reporter_analyse(aid, champs, num, statut, qui)
    sol = next((s for s in a["donnees"]["solutions"] if s["numero"] == num), None)
    registre.ajouter_commentaire({"projet_code": d["projet_code"], "demande_id": d["id"], "a_verifier": True,
                                  "texte": f"Valeurs reportées depuis l'étude RTE n° {a['donnees'].get('num_etude') or '?'}"
                                           + (f", solution {num} ({sol['titre'][:80]})" if sol else "") + " : coûts, délai et numéro à vérifier"}, qui)
    return d


def _reporter_ptf(a: dict, qui: str) -> dict:
    """PTF lue -> demande RTE PTF : numero, contribution, acompte, delais, poste ; « PTF reçue » ; point a verifier."""
    from datetime import date as _date
    x = a["donnees"]
    d = registre.obtenir_demande(a["demande_id"])
    champs = reponses.champs_ptf(x)
    if not d.get("date_ptf_recue"):
        champs["date_ptf_recue"] = _date.today().isoformat()
    statut = "ptf_recue" if d["statut"] in ("pret_validation", "depose", "recevable") else None
    d = registre.reporter_analyse(a["id"], champs, 1, statut, qui)
    fin = None
    if d.get("date_ptf_recue"):
        fin = regles.ajouter_mois(_date.fromisoformat(str(d["date_ptf_recue"])[:10]), int(x.get("validite_mois") or 3)).strftime("%d/%m/%Y")
    registre.ajouter_commentaire({"projet_code": d["projet_code"], "demande_id": d["id"], "a_verifier": True,
                                  "texte": f"Valeurs reportées depuis la PTF n° {x.get('num_ptf') or '?'} : contribution, acompte, délais et poste "
                                           "à vérifier ; date de réception à corriger si la PTF est arrivée avant le dépôt dans Amont"
                                           + (f". Validité jusqu'au {fin}, prorogation de 3 mois à demander avant." if fin else ".")}, qui)
    return d


# ------------------------------------------------- RTE PTF : fiche D2 et courriel d'envoi
def _num_etude(d: dict) -> str | None:
    if d.get("ref_etude_exploratoire"):
        return str(d["ref_etude_exploratoire"])
    for ee in registre.lister_demandes(projet=d["projet_code"], filiere="rte_ee"):
        if ee.get("num_etude"):
            return str(ee["num_etude"])
    return None


def _demande_ptf(id_: str) -> dict:
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    if d["filiere"] != "rte_ptf":
        raise ValueError("la fiche D2 concerne les demandes RTE PTF")
    return d


def _donnees_d2(d: dict) -> dict:
    v = registre.valeurs_fiche(d["projet_code"])
    return d2.donnees(v, d, _num_etude(d), registre.lister_pieces(d["id"]), d["projet_code"])


@app.get("/api/demandes/{id_}/d2")
def apercu_d2(id_: str):
    """Ce que la D2 contiendra, et les cases encore vides."""
    d = _demande_ptf(id_)
    x = _donnees_d2(d)
    return {"donnees": x, "manquants": d2.manquants(x), "elements_complementaires": [t for _, t in d2.elements_complementaires(x) if t],
            "num_etude": x["num_etude"], "pieces_d2": [p for p in registre.lister_pieces(id_) if p["type"] == "fiche_d2"]}


@app.post("/api/demandes/{id_}/d2", status_code=201)
def generer_d2(id_: str, req: Request):
    """Remplit la fiche D2 (.docx) depuis la fiche du site et la joint a la demande (type fiche_d2). Les cases vides
    sortent « à compléter » en jaune, et un point a verifier les liste."""
    from datetime import date as _date
    d = _demande_ptf(id_)
    qui = auteur(req)
    x = _donnees_d2(d)
    contenu = d2.remplir(x)
    nom = f"{_date.today().strftime('%Y_%m_%d')}_fiche_D2_{d['projet_code']}.docx"
    piece = _joindre(id_, "fiche_d2", nom, contenu, qui)
    manque = d2.manquants(x)
    texte = (f"Fiche D2 générée ({piece['nom_fichier']}) : relire, dater, signer, puis joindre le PDF signé (type fiche_d2)."
             + (" Cases à compléter : " + ", ".join(manque) + "." if manque else ""))
    registre.ajouter_commentaire({"projet_code": d["projet_code"], "demande_id": id_, "a_verifier": True, "texte": texte}, qui)
    return {"piece": piece, "manquants": manque, "taille_octets": len(contenu)}


def _contacts_en_copie(d: dict, organisation: str, deja: list[str]) -> list[str]:
    """Adresses des contacts de l'organisation deja en echange sur le projet, sans doublon."""
    cc: list[str] = []
    for e in registre.lister_echanges(projet=d["projet_code"], limite=200):
        c = registre.obtenir_contact(e["contact_id"]) if e.get("contact_id") else None
        if c and c.get("email") and organisation in (c.get("organisation") or "").lower() and c["email"] not in deja + cc:
            cc.append(c["email"])
    return cc


def _message_ptf(d: dict) -> dict:
    v = registre.valeurs_fiche(d["projet_code"])
    a = list(regles.CONFIG.get("rte_ptf_destinataires") or [d2.DESTINATAIRE_RTE])
    cc = _contacts_en_copie(d, "rte", a)
    return d2.courriel(v, d, _num_etude(d), registre.lister_pieces(d["id"]), a, cc)


def _demande_courriel(id_: str) -> dict:
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    if d["filiere"] not in ("rte_ptf", "grdf"):
        raise ValueError("le courriel d'envoi concerne les demandes RTE PTF et GRDF")
    return d


def _message_envoi(d: dict) -> dict:
    return _message_grdf(d) if d["filiere"] == "grdf" else _message_ptf(d)


@app.get("/api/demandes/{id_}/courriel")
def courriel_ptf(id_: str):
    """Brouillon du courriel d'envoi (D2 RTE ou fiche GRDF) : destinataires, objet, corps, pieces retenues, alertes."""
    return _message_envoi(_demande_courriel(id_))


@app.get("/api/demandes/{id_}/courriel.eml")
def courriel_ptf_eml(id_: str):
    """Le meme brouillon en .eml avec les pieces jointes : a ouvrir dans Outlook, relire, envoyer."""
    d = _demande_courriel(id_)
    m = _message_envoi(d)
    v = registre.valeurs_fiche(d["projet_code"])
    fichiers = [(p["nom_fichier"], (DOSSIER_FICHIERS / p["chemin"]).read_bytes()) for p in m["pieces"]]
    nom_exp = " ".join(x for x in (v.get("interlocuteur_prenom"), v.get("interlocuteur_nom")) if x)
    contenu = d2.eml(m, str(v.get("interlocuteur_mail") or ""), nom_exp, fichiers)
    from datetime import date as _date
    nom = f"{_date.today().strftime('%Y_%m_%d')}_courriel_{'GRDF' if d['filiere'] == 'grdf' else 'PTF'}_{d['projet_code']}.eml"
    return StreamingResponse(io.BytesIO(contenu), media_type="message/rfc822",
                             headers={"Content-Disposition": f'attachment; filename="{nom}"'})


# ------------------------------------------------- GRDF : fiche de renseignements Grands Projets
PIECES_PARTAGEES_GAZ = ("plan_situation", "kbis", "plan_200", "plan_2000", "certificat_adressage")


def _demande_grdf(id_: str) -> dict:
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    if d["filiere"] != "grdf":
        raise ValueError("la fiche Grands Projets concerne les demandes GRDF")
    return d


def _pieces_gaz(d: dict) -> list[dict]:
    """Pieces de la demande GRDF, completees par celles des autres demandes du site (plan de situation, cadastre, Kbis...)
    qui servent aussi a GRDF : elles sont deja rangees avec l'etude exploratoire RTE."""
    pieces = registre.lister_pieces(d["id"])
    presents = {p["type"] for p in pieces}
    for autre in registre.lister_demandes(projet=d["projet_code"]):
        if autre["id"] == d["id"]:
            continue
        for p in registre.lister_pieces(autre["id"]):
            if p["type"] in PIECES_PARTAGEES_GAZ and p["type"] not in presents:
                pieces.append(p)
                presents.add(p["type"])
    return pieces


def _donnees_grdf(d: dict) -> dict:
    return grdf.donnees(registre.valeurs_fiche(d["projet_code"]), d, _pieces_gaz(d), d["projet_code"])


@app.get("/api/demandes/{id_}/grdf")
def apercu_grdf(id_: str):
    """Ce que la fiche GRDF contiendra, ses cases encore vides et les points a verifier."""
    d = _demande_grdf(id_)
    x = _donnees_grdf(d)
    plans = [p for p in registre.lister_pieces(id_) if p["type"] == "plan_cadastral_1000"]
    return {"manquants": grdf.manquants(x), "alertes": grdf.alertes_fiche(x),
            "plan_cadastral": plans[-1] if plans else None,
            "pieces_fiche": [p for p in registre.lister_pieces(id_) if p["type"] == "fiche_grdf"],
            "puissance_max_mw": x["puissance_max_mw"], "pression": x["_pression"], "annees": x["_annees"]}


@app.post("/api/demandes/{id_}/grdf", status_code=201)
def generer_grdf(id_: str, req: Request):
    """Remplit la fiche GRDF (PDF) depuis la fiche du site et la joint a la demande (type fiche_grdf). Les champs requis
    vides sortent « à compléter », et un point a verifier les liste."""
    from datetime import date as _date
    d = _demande_grdf(id_)
    qui = auteur(req)
    x = _donnees_grdf(d)
    contenu = grdf.remplir(x)
    nom = f"{_date.today().strftime('%Y_%m_%d')}_fiche_GRDF_{d['projet_code']}.pdf"
    piece = _joindre(id_, "fiche_grdf", nom, contenu, qui)
    manque = grdf.manquants(x)
    texte = (f"Fiche GRDF générée ({piece['nom_fichier']}) : relire, la signer (page 2, nom, date, signature, cachet), puis déposer "
             "le PDF signé (type fiche grdf, nom contenant « signee »)."
             + (" Champs à compléter : " + ", ".join(manque) + "." if manque else "")
             + ((" À vérifier : " + " ; ".join(grdf.alertes_fiche(x)) + ".") if grdf.alertes_fiche(x) else ""))
    registre.ajouter_commentaire({"projet_code": d["projet_code"], "demande_id": id_, "a_verifier": True, "texte": texte}, qui)
    return {"piece": piece, "manquants": manque, "taille_octets": len(contenu)}


def _message_grdf(d: dict) -> dict:
    v = registre.valeurs_fiche(d["projet_code"])
    a = list(regles.CONFIG.get("grdf_destinataires") or [grdf.DESTINATAIRE_GRDF])
    cc = _contacts_en_copie(d, "grdf", a)
    return grdf.courriel(v, d, _pieces_gaz(d), a, cc)


# ------------------------------------------------- robots lances depuis le site
@app.post("/api/demandes/{id_}/robot")
async def lancer_robot(id_: str, req: Request):
    """{"mode": "banc" | "portail", "jusqua": 5} : met une tache en file pour l'agent du poste."""
    corps = await _corps(req)
    t = registre.creer_tache_robot(id_, corps.get("mode") or "banc", corps.get("jusqua") or 5, auteur(req))
    agents = [a for a in registre.agents() if a["en_ligne"]]
    return {**t, "agent_en_ligne": bool(agents)}


@app.get("/api/robot/taches")
def taches_robot(demande: str | None = None, limite: int = 20):
    return {"taches": registre.lister_taches_robot(demande, min(limite, 200)), "agents": registre.agents()}


@app.post("/api/robot/taches/{tid}/annuler")
def annuler_tache_robot(tid: int, req: Request):
    return registre.annuler_tache_robot(tid, auteur(req))


@app.post("/api/robot/prendre")
async def prendre_tache(req: Request):
    """Appele toutes les quelques secondes par l'agent du poste. {"agent": "lenovo-alexis", "version": "..."}"""
    corps = await _corps(req)
    nom = (corps.get("agent") or "poste").strip()[:60]
    return {"tache": registre.prendre_tache_robot(nom, str(corps.get("version") or ""), str(corps.get("info") or ""))}


@app.post("/api/robot/taches/{tid}/journal")
async def journal_tache(tid: int, req: Request):
    corps = await _corps(req)
    registre.journal_tache_robot(tid, [str(x)[:2000] for x in corps.get("lignes") or []][:500])
    return {"ok": True}


@app.post("/api/robot/taches/{tid}/fin")
async def fin_tache(tid: int, req: Request):
    corps = await _corps(req)
    return registre.finir_tache_robot(tid, corps.get("statut") or "terminee", str(corps.get("resultat") or "")[:500])


@app.post("/api/demandes/{id_}/plan-installation")
def plan_installation(id_: str, req: Request):
    return assembler_plan_installation(id_, auteur(req))


@app.patch("/api/sites/{code}/fiche")
async def modifier_fiche(code: str, req: Request):
    return registre.modifier_fiche(code, await req.json(), auteur(req))


@app.post("/api/sites/{code}/fiche/valider")
async def valider_fiche(code: str, req: Request):
    corps = await req.json()
    champs = corps.get("champs") or [k for k, v in registre.fiche(code).items() if not v["valide"]]
    return registre.valider_champs(code, champs, auteur(req))


@app.post("/api/sites/{code}/fiche/import")
async def importer_fiche(code: str, req: Request):
    """Entree pour un autre programme (DatacenterSiteAnalyzer...) : {"origine": "dcsa", "valeurs": {champ: valeur}}."""
    corps = await req.json()
    return registre.importer_fiche(code, corps.get("valeurs") or {}, corps.get("origine") or "import", auteur(req),
                                   bool(corps.get("ecraser")))


@app.post("/api/propositions/{prop_id}/accepter")
def accepter(prop_id: int, req: Request):
    return registre.accepter_proposition(prop_id, auteur(req))


@app.post("/api/propositions/{prop_id}/rejeter")
def rejeter(prop_id: int, req: Request):
    return registre.rejeter_proposition(prop_id, auteur(req))


def _extraire_source(source_id: int, qui: str) -> None:
    s = registre.obtenir_source(source_id)
    if not s:
        return
    registre.marquer_source(source_id, statut="en_cours", erreur=None)
    try:
        res = extraction.extraire(DOSSIER_SOURCES / s["chemin"], s["nom_fichier"], registre.valeurs_fiche(s["projet_code"]))
        registre.enregistrer_extraction(source_id, res, qui)
    except Exception as e:  # ne jamais laisser une source bloquee « en cours »
        registre.marquer_source(source_id, statut="erreur", erreur=f"{type(e).__name__} : {e}")


@app.post("/api/sites/{code}/sources", status_code=201)
async def deposer_sources(code: str, req: Request, taches: BackgroundTasks, fichiers: list[UploadFile] = File(...)):
    if not registre.obtenir_projet(code):
        raise KeyError(code)
    qui = auteur(req)
    dossier = DOSSIER_SOURCES / code
    dossier.mkdir(parents=True, exist_ok=True)
    crees = []
    for f in fichiers:
        nom = re.sub(r'[\\/:*?"<>|]', "_", Path(f.filename or "source").name)
        if Path(nom).suffix.lower() not in EXTENSIONS_SOURCES:
            raise ValueError(f"format non pris en charge : {nom}")
        contenu = await f.read()
        if len(contenu) > 60 * 1024 * 1024:
            raise ValueError(f"fichier au-delà de 60 Mo : {nom}")
        cible = dossier / nom
        i = 1
        while cible.exists():
            cible = dossier / f"{Path(nom).stem}_{i}{Path(nom).suffix}"
            i += 1
        cible.write_bytes(contenu)
        s = registre.ajouter_source(code, cible.name, str(cible.relative_to(DOSSIER_SOURCES)).replace("\\", "/"), len(contenu), qui)
        taches.add_task(_extraire_source, s["id"], qui)
        crees.append(s)
    return crees


@app.post("/api/sources/{source_id}/extraire")
def reextraire(source_id: int, req: Request, taches: BackgroundTasks):
    if not registre.obtenir_source(source_id):
        raise KeyError(source_id)
    registre.marquer_source(source_id, statut="en_attente")
    taches.add_task(_extraire_source, source_id, auteur(req))
    return {"ok": True}


@app.post("/api/sites/{code}/extraire")
def extraire_tout(code: str, req: Request, taches: BackgroundTasks):
    ids = [s["id"] for s in registre.lister_sources(code)]
    for i in ids:
        registre.marquer_source(i, statut="en_attente")
        taches.add_task(_extraire_source, i, auteur(req))
    return {"sources": len(ids)}


@app.get("/api/sources/{source_id}/fichier")
def fichier_source(source_id: int):
    s = registre.obtenir_source(source_id)
    if not s:
        raise KeyError(source_id)
    chemin = DOSSIER_SOURCES / s["chemin"]
    if not chemin.exists():
        raise HTTPException(404, "fichier absent du disque")
    return FileResponse(chemin, filename=s["nom_fichier"])


@app.delete("/api/sources/{source_id}", status_code=204)
def supprimer_source(source_id: int, req: Request):
    s = registre.supprimer_source(source_id, auteur(req))
    if not s:
        raise KeyError(source_id)
    chemin = DOSSIER_SOURCES / s["chemin"]
    if chemin.exists():
        chemin.unlink()


@app.post("/api/sites/{code}/preparer")
async def preparer(code: str, req: Request):
    """Cree d'un coup les demandes absentes du site, pre-remplies depuis la fiche."""
    corps = await req.json()
    return registre.preparer_demandes(code, corps.get("filieres") or regles.ORDRE_FILIERES, auteur(req))


@app.post("/api/demandes/{id_}/synchroniser")
def synchroniser(id_: str, req: Request):
    return registre.synchroniser_demande(id_, auteur(req))

# ---------------------------------------------------------------------- pieces
@app.post("/api/demandes/{id_}/pieces", status_code=201)
async def ajouter_piece(id_: str, req: Request, fichier: UploadFile = File(...), type: str = Form("autre")):
    d = registre.obtenir_demande(id_)
    if not d:
        raise KeyError(id_)
    contenu = await fichier.read()
    if len(contenu) > 25 * 1024 * 1024:
        raise ValueError("fichier au-delà de 25 Mo")
    nom = re.sub(r"[^A-Za-z0-9_.\-]", "_", fichier.filename or "piece.pdf")
    dossier = DOSSIER_FICHIERS / d["projet_code"] / d["filiere"]
    dossier.mkdir(parents=True, exist_ok=True)
    cible = dossier / nom
    i = 1
    while cible.exists():
        cible = dossier / f"{cible.stem.rsplit('_', 1)[0] if i > 1 else cible.stem}_{i}{cible.suffix}"
        i += 1
    cible.write_bytes(contenu)
    chemin = str(cible.relative_to(DOSSIER_FICHIERS)).replace("\\", "/")
    return registre.ajouter_piece(id_, type, cible.name, chemin, len(contenu), auteur(req))


@app.get("/api/pieces/{piece_id}/fichier")
def fichier_piece(piece_id: int):
    r = registre.cx.execute("SELECT * FROM pieces WHERE id = ?", (piece_id,)).fetchone()
    if not r:
        raise KeyError(piece_id)
    chemin = DOSSIER_FICHIERS / r["chemin"]
    if not chemin.exists():
        raise HTTPException(404, "fichier absent du disque")
    return FileResponse(chemin, filename=r["nom_fichier"])


@app.delete("/api/pieces/{piece_id}", status_code=204)
def supprimer_piece(piece_id: int, req: Request):
    r = registre.supprimer_piece(piece_id, auteur(req))
    if not r:
        raise KeyError(piece_id)
    chemin = DOSSIER_FICHIERS / r["chemin"]
    if chemin.exists():
        chemin.unlink()


# -------------------------------------------------- commentaires et verifications
@app.get("/api/commentaires")
def commentaires(projet: str | None = None, demande: str | None = None, a_verifier: bool = False, ouverts: bool = False):
    return registre.lister_commentaires(projet, demande, a_verifier, ouverts)


@app.post("/api/commentaires", status_code=201)
async def creer_commentaire(req: Request):
    """{projet_code ou demande_id, texte, a_verifier, cible} ; cible « fiche:<champ> » signale un champ de la fiche."""
    return registre.ajouter_commentaire(await req.json(), auteur(req))


@app.patch("/api/commentaires/{cid}")
async def traiter_commentaire(cid: int, req: Request):
    corps = await req.json()
    return registre.traiter_commentaire(cid, corps.get("statut") or "verifie", corps.get("reponse") or "", auteur(req))


@app.delete("/api/commentaires/{cid}", status_code=204)
def supprimer_commentaire(cid: int, req: Request):
    registre.supprimer_commentaire(cid, auteur(req))


# ------------------------------------------------------------------ ressources
DOSSIER_RESSOURCES = DOSSIER_FICHIERS.parent / "ressources"


@app.get("/api/ressources")
def ressources(filiere: str | None = None, etape: str | None = None):
    return registre.lister_ressources(filiere, etape)


@app.post("/api/ressources", status_code=201)
async def creer_ressource(req: Request):
    return registre.creer_ressource(await req.json(), auteur(req))


@app.patch("/api/ressources/{rid}")
async def modifier_ressource(rid: int, req: Request):
    return registre.modifier_ressource(rid, await req.json(), auteur(req))


@app.delete("/api/ressources/{rid}", status_code=204)
def supprimer_ressource(rid: int, req: Request):
    r = registre.supprimer_ressource(rid, auteur(req))
    if not r:
        raise KeyError(rid)
    if r.get("fichier") and (DOSSIER_RESSOURCES / r["fichier"]).exists():
        (DOSSIER_RESSOURCES / r["fichier"]).unlink()


@app.post("/api/ressources/{rid}/fichier")
async def deposer_ressource(rid: int, req: Request, fichier: UploadFile = File(...)):
    r = registre.obtenir_ressource(rid)
    if not r:
        raise KeyError(rid)
    contenu = await fichier.read()
    if len(contenu) > 25 * 1024 * 1024:
        raise ValueError("fichier au-delà de 25 Mo")
    nom = re.sub(r"[^A-Za-z0-9_.\-]", "_", Path(fichier.filename or "document").name)
    DOSSIER_RESSOURCES.mkdir(parents=True, exist_ok=True)
    cible = DOSSIER_RESSOURCES / f"{rid}_{nom}"
    if r.get("fichier") and (DOSSIER_RESSOURCES / r["fichier"]).exists():
        (DOSSIER_RESSOURCES / r["fichier"]).unlink()
    cible.write_bytes(contenu)
    return registre.modifier_ressource(rid, {"fichier": cible.name, "nom_fichier": nom}, auteur(req))


@app.get("/api/ressources/{rid}/fichier")
def fichier_ressource(rid: int):
    r = registre.obtenir_ressource(rid)
    if not r or not r.get("fichier"):
        raise KeyError(rid)
    chemin = DOSSIER_RESSOURCES / r["fichier"]
    if not chemin.exists():
        raise HTTPException(404, "fichier absent du disque")
    return FileResponse(chemin, filename=r["nom_fichier"] or chemin.name)


# -------------------------------------------------------------------- contacts
@app.get("/api/contacts")
def contacts():
    return registre.lister_contacts()


@app.post("/api/contacts", status_code=201)
async def creer_contact(req: Request):
    return registre.creer_contact(await req.json(), auteur(req))


@app.patch("/api/contacts/{contact_id}")
async def modifier_contact(contact_id: int, req: Request):
    return registre.modifier_contact(contact_id, await req.json(), auteur(req))


@app.delete("/api/contacts/{contact_id}", status_code=204)
def supprimer_contact(contact_id: int, req: Request):
    registre.supprimer_contact(contact_id, auteur(req))


# -------------------------------------------------------------------- echanges
@app.get("/api/echanges")
def echanges_liste(projet: str | None = None, demande: str | None = None, attente: bool = False,
                   contact: int | None = None, limite: int = 200):
    return registre.lister_echanges(projet, demande, attente, contact, min(limite, 2000))


@app.post("/api/echanges", status_code=201)
async def creer_echange(req: Request):
    """Saisie manuelle depuis le site, et point d'entree du futur releve de boite mail.

    Corps JSON : projet_code ou demande_id, sens (envoye, recu, note), canal (mail...), date,
    expediteur, destinataires, objet, resume, attend_reponse, relance_le, contact_id, message_id.
    Un programme de releve passe l'en-tete X-Amont-Origine (ex. « graph ») et le message_id,
    qui sert au dedoublonnage : renvoyer deux fois le meme courriel ne cree qu'un echange.
    """
    corps = await req.json()
    origine = req.headers.get("X-Amont-Origine") or corps.pop("origine", None) or "manuel"
    if not corps.get("demande_id") and not corps.get("projet_code"):
        texte = " ".join(str(corps.get(k) or "") for k in ("objet", "resume"))
        corps["demande_id"], corps["projet_code"] = courriels.deviner_demande(texte, registre.lister_demandes())
    return registre.ajouter_echange(corps, auteur(req), origine)


@app.post("/api/echanges/eml", status_code=201)
async def deposer_courriels(req: Request, fichiers: list[UploadFile] = File(...), projet_code: str = Form(""),
                            demande_id: str = Form("")):
    """Depot d'un ou plusieurs .eml. Le projet et la demande sont devines si non fournis ;
    un courriel non rattachable est refuse avec son objet, pour etre redepose depuis la bonne fiche."""
    qui = auteur(req)
    demandes = registre.lister_demandes()
    crees, refuses = [], []
    for f in fichiers:
        nom = Path(f.filename or "courriel.eml").name
        if Path(nom).suffix.lower() != ".eml":
            refuses.append({"fichier": nom, "motif": "format .eml attendu (Outlook : Télécharger le message, ou glisser le mail sur le bureau)"})
            continue
        contenu = await f.read()
        if len(contenu) > TAILLE_MAX_COURRIEL:
            refuses.append({"fichier": nom, "motif": "au-delà de 30 Mo"})
            continue
        m = courriels.lire_eml(contenu)
        dem, proj = demande_id or None, projet_code or None
        if not dem:
            texte = m["objet"] + "\n" + m["corps"]
            dem_g, proj_g = courriels.deviner_demande(texte, [d for d in demandes if not proj or d["projet_code"] == proj])
            dem = dem_g
            proj = proj or proj_g
            if not dem and proj:
                fil = courriels.deviner_filiere(texte, m["expediteur"] + " " + m["destinataires"])
                cand = [d for d in demandes if d["projet_code"] == proj and d["filiere"] == fil]
                if len(cand) == 1:
                    dem = cand[0]["id"]
        if dem:
            proj = next((d["projet_code"] for d in demandes if d["id"] == dem), proj)
        if not proj:
            refuses.append({"fichier": nom, "objet": m["objet"], "motif": "projet introuvable, déposer depuis la fiche du projet ou de la demande"})
            continue
        dossier = DOSSIER_FICHIERS / proj / "courriels"
        dossier.mkdir(parents=True, exist_ok=True)
        jour = (m["date"] or courriels.maintenant())[:10].replace("-", "_")
        base = re.sub(r"[^A-Za-z0-9_.\-]", "_", f"{jour}_{m['objet'][:60] or 'courriel'}").strip("_")
        cible = dossier / f"{base}.eml"
        i = 1
        while cible.exists():
            cible = dossier / f"{base}_{i}.eml"
            i += 1
        sens = courriels.deviner_sens(m["expediteur"], _mes_adresses())
        e = registre.ajouter_echange({
            "projet_code": proj, "demande_id": dem, "sens": sens, "canal": "mail", "date": m["date"],
            "expediteur": m["expediteur"], "destinataires": m["destinataires"], "objet": m["objet"],
            "resume": courriels.resume(m["corps"]) + (("\n\nPièces jointes : " + ", ".join(m["pieces_jointes"])) if m["pieces_jointes"] else ""),
            "message_id": m["message_id"], "nom_fichier": nom,
            "fichier": str(cible.relative_to(DOSSIER_FICHIERS)).replace("\\", "/"),
        }, qui, "eml")
        if not e.get("doublon"):
            cible.write_bytes(contenu)
            # etude RTE en piece jointe d'un courriel recu : rangee dans la demande et lue
            if dem and sens == "recu":
                for nom_pj, data in courriels.pieces_jointes(contenu):
                    try:
                        import tempfile
                        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
                            tmp.write(data)
                        est_etude = reponses.est_etude_rte(_texte_pdf(Path(tmp.name)))
                        Path(tmp.name).unlink()
                        if est_etude:
                            pj = _joindre(dem, "etude_recue", re.sub(r"[^A-Za-z0-9_.\-]", "_", nom_pj), data, qui)
                            e.setdefault("etudes_lues", []).append(lire_piece(pj["id"], qui)["id"])
                    except Exception:  # noqa: BLE001 : une piece jointe illisible ne bloque pas le courriel
                        pass
        crees.append(e)
    return {"crees": crees, "refuses": refuses}


@app.patch("/api/echanges/{echange_id}")
async def modifier_echange(echange_id: int, req: Request):
    return registre.modifier_echange(echange_id, await req.json(), auteur(req))


@app.delete("/api/echanges/{echange_id}", status_code=204)
def supprimer_echange(echange_id: int, req: Request):
    e = registre.supprimer_echange(echange_id, auteur(req))
    if not e:
        raise KeyError(echange_id)
    if e.get("fichier"):
        chemin = DOSSIER_FICHIERS / e["fichier"]
        if chemin.exists():
            chemin.unlink()


@app.get("/api/echanges/{echange_id}/fichier")
def fichier_echange(echange_id: int):
    e = registre.obtenir_echange(echange_id)
    if not e or not e.get("fichier"):
        raise KeyError(echange_id)
    chemin = DOSSIER_FICHIERS / e["fichier"]
    if not chemin.exists():
        raise HTTPException(404, "fichier absent du disque")
    return FileResponse(chemin, filename=e["nom_fichier"] or chemin.name, media_type="message/rfc822")


# -------------------------------------------------------- echeances et journal
@app.get("/api/echeances")
def echeances(severite: str | None = None, avant: str | None = None):
    out = []
    for d in registre.lister_demandes():
        for e in d["echeances"]:
            if e["fait"]:
                continue
            if severite and e["severite"] != severite:
                continue
            if avant and e["date"] > avant:
                continue
            out.append({**e, "demande_id": d["id"], "projet": d["projet_code"], "filiere": d["filiere"]})
    out.sort(key=lambda e: e["date"])
    return out


@app.get("/api/alertes")
def toutes_alertes():
    out = []
    for d in registre.lister_demandes():
        for a in d["alertes"]:
            out.append({**a, "demande_id": d["id"], "projet": d["projet_code"], "filiere": d["filiere"], "site": d.get("site")})
    out.sort(key=lambda a: 0 if a["severite"] == "crit" else 1)
    return out


@app.get("/api/alerte/apercu")
def apercu_courriel():
    return alertes.apercu(registre, URL_BASE)


@app.get("/api/journal")
def journal(limite: int = 50):
    return registre.journal(limite=min(limite, 1000))


# ---------------------------------------------------------------------- export
@app.get("/api/export.xlsx")
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    projets = {p["code"]: p["nom"] for p in registre.lister_projets()}
    wb = Workbook()
    ws = wb.active
    ws.title = "Demandes"
    entete = ["Projet", "Code", "Filière", "Site", "Commune", "Dépt", "Étape", "Dimension", "Date de dépôt",
              "Prochaine échéance", "Date", "Jours", "Prochaine action", "Dernier échange", "Sans réponse depuis (j)",
              "Coût client k€", "Interlocuteur", "Modifié le", "Par"]
    ws.append(entete)
    for c in ws[1]:
        c.font = Font(bold=True)
    for d in registre.lister_demandes():
        pe = d.get("prochaine_echeance") or {}
        ws.append([projets.get(d["projet_code"], d["projet_code"]), d["id"], regles.FILIERES[d["filiere"]]["court"],
                   d.get("site"), d.get("commune"), d.get("departement"), d["libelle_statut"], d["dimension"],
                   d.get("date_depot"), pe.get("libelle"), pe.get("date"), pe.get("jours"),
                   d["action"]["texte"], ((d["echanges_resume"]["dernier"] or {}).get("date") or "")[:10],
                   (d["echanges_resume"]["attente"] or {}).get("jours"),
                   d.get("cout_client_keur"), d.get("interlocuteur"), d.get("modifie_le"), d.get("modifie_par")])
    for i, largeur in enumerate([16, 16, 9, 28, 16, 6, 22, 12, 12, 44, 12, 7, 50, 14, 10, 14, 22, 18, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Échéances")
    ws2.append(["Date", "Jours", "Sévérité", "Demande", "Projet", "Filière", "Échéance"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for e in echeances():
        ws2.append([e["date"], e["jours"], e["severite"], e["demande_id"], projets.get(e["projet"], e["projet"]),
                    regles.FILIERES[e["filiere"]]["court"], e["libelle"]])
    for i, largeur in enumerate([12, 7, 9, 16, 16, 9, 60], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = largeur
    ws3 = wb.create_sheet("Échanges")
    ws3.append(["Date", "Sens", "Canal", "Projet", "Demande", "Contact", "Objet", "Attend une réponse", "Relance le", "Résumé"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for e in registre.lister_echanges(limite=5000):
        ws3.append([e["date"], e["libelle_sens"], e["libelle_canal"], projets.get(e["projet_code"], e["projet_code"]),
                    e["demande_id"], e["contact_nom"] or e["expediteur"] or e["destinataires"], e["objet"],
                    "oui" if e["attend_reponse"] else "", e["relance_le"], (e["resume"] or "")[:500]])
    for i, largeur in enumerate([17, 8, 11, 16, 16, 26, 44, 10, 11, 80], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = largeur
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from datetime import date
    nom = f"{date.today().strftime('%Y_%m_%d')}_amont_demandes.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nom}"'})


# ----------------------------------------------------------------------- front
app.mount("/", StaticFiles(directory=str(RACINE / "static"), html=True), name="static")
