# -*- coding: utf-8 -*-
"""
DCEX - Générateur automatique de pré-études exploratoires de raccordement RTE
pour projets datacenter, au format des études exploratoires RTE.

Usage :
    python 2026_07_13_dcex_etude_explo.py --site site.json
    python 2026_07_13_dcex_etude_explo.py --site site.json --offline cache.json

Dépendances :
    pip install requests python-docx matplotlib

Le fichier site.json décrit la demande (voir exemple fourni).
Le document Word généré reprend la structure d'une étude exploratoire RTE
(préambule, hypothèses, solution de raccordement, notas, conclusion).
ATTENTION : les résultats d'étude de transit (contraintes réseau,
renforcements, limitations temporaires) ne sont PAS calculables sans le
modèle de réseau interne RTE : ils apparaissent comme champs [À COMPLÉTER]
ou estimations paramétrables. Ce document est une pré-étude interne
indicative, il n'est pas émis par RTE.
"""

import argparse
import json
import math
import os
import re
import sys
import unicodedata
from datetime import date

try:
    import requests
except ImportError:
    requests = None

# ---------------------------------------------------------------------------
# Constantes API
# ---------------------------------------------------------------------------

GEO_API = "https://geo.api.gouv.fr/communes"
ODRE_LEGACY = "https://odre.opendatasoft.com/api/records/1.0/search/"
ODRE_V21 = "https://odre.opendatasoft.com/api/explore/v2.1/catalog/datasets/{ds}/records"

DATASETS = {
    "lignes_aeriennes": ["lignes-aeriennes-rte-nv", "lignes-aeriennes-rte"],
    "lignes_souterraines": ["lignes-souterraines-rte-nv", "lignes-souterraines-rte"],
    "postes": ["postes-electriques-rte-et-tiers", "postes-electriques-rte"],
}

# Tableau de l'article 105 de l'arrêté du 9 juin 2020 (simplifié)
DOMAINES_TENSION = [
    ("HTA", 40, "20 kV"),
    ("HTB1", 100, "63 kV / 90 kV"),
    ("HTB2", 400, "150 kV / 225 kV"),
    ("HTB3", float("inf"), "400 kV"),
]

TENSIONS_KV = {"HTB1": [63, 90], "HTB2": [150, 225], "HTB3": [400]}

# Coefficient K du critere K/d de l'article 105
COEF_K = {"HTA": 100, "HTB1": 1000, "HTB2": 10000}


# ---------------------------------------------------------------------------
# Utilitaires
# ---------------------------------------------------------------------------

def haversine_km(lat1, lon1, lat2, lon2):
    r = 6371.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s)
                   if unicodedata.category(c) != "Mn")


def extraire_tension_kv(record):
    """Retourne la tension en kV d'un enregistrement ODRE (champs variables)."""
    for key, val in record.items():
        if "tension" in key.lower() and val is not None:
            m = re.search(r"(\d+)", str(val))
            if m:
                v = int(m.group(1))
                if v in (400, 225, 150, 90, 63, 45, 42, 20):
                    return v
    return None


def extraire_noms(record):
    """Concatène les champs nom_ligne* / nom_poste* / nom* d'un record ODRE."""
    noms = []
    for key in sorted(record.keys()):
        kl = key.lower()
        if kl.startswith("nom") and record[key]:
            v = str(record[key]).strip()
            if v and v not in noms:
                noms.append(v)
    return " / ".join(noms) if noms else "(nom non renseigné)"


def geometrie_points(geometry):
    """Extrait la liste [(lat, lon)] d'une géométrie GeoJSON."""
    pts = []
    if not geometry:
        return pts
    t = geometry.get("type")
    c = geometry.get("coordinates")
    if t == "Point":
        pts.append((c[1], c[0]))
    elif t == "LineString":
        pts.extend((p[1], p[0]) for p in c)
    elif t == "MultiLineString":
        for line in c:
            pts.extend((p[1], p[0]) for p in line)
    return pts


# ---------------------------------------------------------------------------
# Acquisition des données
# ---------------------------------------------------------------------------

def geocoder_commune(nom_commune, code_dept=None):
    params = {"nom": nom_commune, "fields": "nom,code,centre,codeDepartement",
              "boost": "population"}
    r = requests.get(GEO_API, params=params, timeout=30)
    r.raise_for_status()
    communes = r.json()
    if code_dept:
        communes = [c for c in communes
                    if c.get("codeDepartement") == str(code_dept)] or communes
    if not communes:
        raise ValueError(f"Commune introuvable : {nom_commune}")
    c = communes[0]
    lon, lat = c["centre"]["coordinates"]
    return {"nom": c["nom"], "insee": c["code"],
            "departement": c.get("codeDepartement"), "lat": lat, "lon": lon}


def requete_odre(dataset, lat, lon, rayon_km, rows=200):
    """Interroge ODRE (API legacy puis v2.1 en secours)."""
    # API legacy
    try:
        params = {"dataset": dataset, "rows": rows,
                  "geofilter.distance": f"{lat},{lon},{int(rayon_km * 1000)}"}
        r = requests.get(ODRE_LEGACY, params=params, timeout=60)
        if r.status_code == 200:
            data = r.json()
            if "records" in data:
                out = []
                for rec in data["records"]:
                    fields = rec.get("fields", {})
                    geom = rec.get("geometry") or fields.get("geo_shape")
                    if isinstance(geom, dict) and geom.get("type") == "Feature":
                        geom = geom.get("geometry")
                    out.append({"fields": fields, "geometry": geom})
                return out
    except Exception as e:
        print(f"  [!] API legacy KO pour {dataset} : {e}")
    # API v2.1
    for geocol in ("geo_shape", "geo_point_2d"):
        try:
            where = (f"within_distance({geocol}, geom'POINT({lon} {lat})', "
                     f"{int(rayon_km)}km)")
            r = requests.get(ODRE_V21.format(ds=dataset),
                             params={"where": where, "limit": min(rows, 100)},
                             timeout=60)
            if r.status_code == 200:
                data = r.json()
                out = []
                for rec in data.get("results", []):
                    geom = rec.get("geo_shape")
                    if isinstance(geom, dict) and geom.get("type") == "Feature":
                        geom = geom.get("geometry")
                    out.append({"fields": rec, "geometry": geom})
                return out
        except Exception as e:
            print(f"  [!] API v2.1 ({geocol}) KO pour {dataset} : {e}")
    return []


def collecter_reseau(lat, lon, rayon_km, offline_cache=None):
    """Collecte lignes et postes RTE autour du site. Renvoie un dict brut."""
    if offline_cache:
        with open(offline_cache, encoding="utf-8") as f:
            return json.load(f)
    if requests is None:
        raise RuntimeError("Le module requests est requis (pip install requests)")
    reseau = {}
    for cle, candidats in DATASETS.items():
        records = []
        for ds in candidats:
            print(f"  Interrogation ODRE : {ds} ...")
            records = requete_odre(ds, lat, lon, rayon_km)
            if records:
                print(f"    -> {len(records)} enregistrements")
                break
        reseau[cle] = records
    return reseau


# ---------------------------------------------------------------------------
# Analyse
# ---------------------------------------------------------------------------

def tension_de_reference(puissance_mw, d_km=None):
    """Domaine de tension de reference selon l'article 105 de l'arrete du
    9 juin 2020.

    La regle complete est : Prac doit etre inferieure a la PLUS PETITE des
    deux valeurs, un seuil absolu et K/d, ou d est la distance en km sur un
    parcours de reseau realisable entre le point de raccordement et le point
    de transformation vers la tension superieure le plus proche.

    Sans d, seul le seuil absolu est applique (resultat optimiste : la
    tension reelle peut etre superieure).
    Valide sur 5 solutions RTE reelles (Angouleme, Haynecourt x3, Fortum).
    """
    for domaine, seuil, libelle in DOMAINES_TENSION:
        if domaine == "HTB3":
            break
        limite = seuil if d_km in (None, 0) else min(seuil, COEF_K[domaine] / d_km)
        if puissance_mw < limite:
            return domaine, libelle
    return "HTB3", "400 kV"


def puissance_max_admissible(domaine, d_km):
    """Puissance maximale raccordable a un domaine de tension donne compte
    tenu de d. Sert a proposer une solution degradee (cf. etude Fortum :
    64 MW imposent le 225 kV, mais 40 MW permettent le 63 kV)."""
    seuil = dict((d, s) for d, s, _ in DOMAINES_TENSION)[domaine]
    if d_km in (None, 0):
        return seuil
    return min(seuil, COEF_K[domaine] / d_km)


def distance_point_transformation(ouvrages, domaine, coef_sinuosite=1.25):
    """Estime d : distance au poste le plus proche assurant la transformation
    vers la tension immediatement superieure. Distance a vol d'oiseau
    majoree d'un coefficient de sinuosite, RTE retenant un parcours reseau
    realisable techniquement et administrativement."""
    superieurs = {"HTA": [63, 90, 150, 225, 400],
                  "HTB1": [150, 225, 400],
                  "HTB2": [400]}.get(domaine, [400])
    candidats = [p for p in ouvrages["postes"]
                 if p["tension_kv"] in superieurs]
    if not candidats:
        return None, None
    poste = min(candidats, key=lambda p: p["distance_km"])
    return round(poste["distance_km"] * coef_sinuosite, 1), poste


def analyser_ouvrages(reseau, lat, lon):
    """Calcule pour chaque ouvrage sa distance au site et sa tension."""
    ouvrages = {"lignes": [], "postes": []}
    for cle in ("lignes_aeriennes", "lignes_souterraines"):
        for rec in reseau.get(cle, []):
            pts = geometrie_points(rec.get("geometry"))
            if not pts:
                gp = rec["fields"].get("geo_point_2d")
                if isinstance(gp, (list, tuple)):
                    pts = [(gp[0], gp[1])]
                elif isinstance(gp, dict):
                    pts = [(gp.get("lat"), gp.get("lon"))]
            if not pts:
                continue
            d = min(haversine_km(lat, lon, p[0], p[1]) for p in pts)
            ouvrages["lignes"].append({
                "nom": extraire_noms(rec["fields"]),
                "tension_kv": extraire_tension_kv(rec["fields"]),
                "distance_km": round(d, 2),
                "type": ("aérienne" if cle == "lignes_aeriennes"
                         else "souterraine"),
                "points": pts,
            })
    for rec in reseau.get("postes", []):
        pts = geometrie_points(rec.get("geometry"))
        if not pts:
            gp = rec["fields"].get("geo_point_2d")
            if isinstance(gp, (list, tuple)):
                pts = [(gp[0], gp[1])]
            elif isinstance(gp, dict):
                pts = [(gp.get("lat"), gp.get("lon"))]
        if not pts:
            continue
        d = min(haversine_km(lat, lon, p[0], p[1]) for p in pts)
        ouvrages["postes"].append({
            "nom": extraire_noms(rec["fields"]),
            "tension_kv": extraire_tension_kv(rec["fields"]),
            "distance_km": round(d, 2),
            "points": pts,
        })
    ouvrages["lignes"].sort(key=lambda o: o["distance_km"])
    ouvrages["postes"].sort(key=lambda o: o["distance_km"])
    return ouvrages


def construire_solution(site, ouvrages, config):
    """Détermine la solution de raccordement candidate et son coût indicatif."""
    p_max = site["puissance_max_mw"]
    coef_sin = config.get("coefficient_sinuosite", 1.25)
    # Critere d : distance au point de transformation vers la tension
    # superieure. Iteration car d depend du domaine teste.
    domaine, libelle_tension = tension_de_reference(p_max)
    d_km, poste_transfo = distance_point_transformation(
        ouvrages, domaine, coef_sin)
    if d_km:
        domaine, libelle_tension = tension_de_reference(p_max, d_km)
        d_km2, poste_transfo2 = distance_point_transformation(
            ouvrages, domaine, coef_sin)
        if d_km2:
            d_km, poste_transfo = d_km2, poste_transfo2
    if domaine == "HTA":
        raise ValueError("Puissance < 40 MW : raccordement HTA (Enedis), hors "
                         "périmètre de cet outil.")
    tensions_cibles = TENSIONS_KV[domaine]

    lignes_ok = [l for l in ouvrages["lignes"]
                 if l["tension_kv"] in tensions_cibles]
    postes_ok = [p for p in ouvrages["postes"]
                 if p["tension_kv"] in tensions_cibles]

    ligne = lignes_ok[0] if lignes_ok else None
    poste = postes_ok[0] if postes_ok else None

    couts = config["couts_unitaires_keur"][domaine]
    seuil_poste = config["seuil_raccordement_poste_km"]

    solution = {"domaine": domaine, "tension": libelle_tension,
                "ligne_proche": ligne, "poste_proche": poste,
                "d_km": d_km,
                "poste_transformation": (poste_transfo["nom"]
                                         if poste_transfo else None)}
    # Solution degradee : puissance plafond au domaine de tension inferieur
    inferieur = {"HTB3": "HTB2", "HTB2": "HTB1", "HTB1": "HTA"}.get(domaine)
    if inferieur and inferieur != "HTA":
        solution["puissance_max_domaine_inferieur"] = round(
            puissance_max_admissible(inferieur, d_km), 1)

    if poste and (not ligne or poste["distance_km"] <= ligne["distance_km"]
                  or poste["distance_km"] <= seuil_poste):
        # Raccordement en antenne sur poste existant
        nb_cell = config["nb_cellules_defaut"]["antenne_poste_existant"]
        long_liaison = max(poste["distance_km"], 0.1)
        strategie = (f"Raccordement en antenne sur le poste existant "
                     f"{poste['nom']} moyennant la création d'une liaison "
                     f"{libelle_tension} d'environ {long_liaison:.1f} km")
        cout = (nb_cell * couts["cellule_ligne"]
                + long_liaison * couts[config.get("type_liaison_defaut",
                                 "km_liaison_souterraine")])
        travaux = [
            f"Création de {nb_cell} cellules {libelle_tension} au poste "
            f"{poste['nom']}.",
            f"Création d'une liaison {libelle_tension} d'environ "
            f"{long_liaison:.1f} km entre le poste {poste['nom']} et le point "
            f"de raccordement.",
        ]
        solution["type"] = "antenne_poste_existant"
    elif ligne:
        # Entrée en coupure sur ligne(s) existante(s) + nouveau poste
        double = " n° 1 et n° 2" in ligne["nom"] or "N.1" in ligne["nom"].upper()
        cle_cell = "double_coupure" if double else "simple_coupure"
        nb_cell = config["nb_cellules_defaut"][cle_cell]
        long_liaison = max(ligne["distance_km"], 0.1)
        mode = ("double entrée en coupure" if double else "entrée en coupure")
        strategie = (f"Raccordement en {mode} sur la ligne {libelle_tension} "
                     f"{ligne['nom']} moyennant la création d'un nouveau poste "
                     f"et d'une liaison {libelle_tension} d'environ "
                     f"{long_liaison:.1f} km entre ce poste et le point de "
                     f"raccordement")
        cout = (nb_cell * couts["cellule_ligne"]
                + couts["genie_civil_poste"]
                + long_liaison * couts[config.get("type_liaison_defaut",
                                 "km_liaison_souterraine")])
        travaux = [
            f"Création d'un poste en {mode} sur la ligne {libelle_tension} "
            f"{ligne['nom']} comprenant {nb_cell} cellules lignes.",
            f"Création d'une liaison souterraine {libelle_tension} d'environ "
            f"{long_liaison:.1f} km entre le nouveau poste et le point de "
            f"raccordement.",
        ]
        solution["type"] = "coupure_ligne"
    else:
        strategie = (f"Aucun ouvrage {libelle_tension} identifié dans le rayon "
                     f"de recherche. Raccordement nécessitant un développement "
                     f"de réseau important [À COMPLÉTER].")
        cout = 0
        nb_cell = 0
        long_liaison = 0
        travaux = ["[À COMPLÉTER] : aucun ouvrage au niveau de tension de "
                   "référence dans le rayon étudié."]
        solution["type"] = "hors_rayon"

    solution.update({
        "strategie": strategie,
        "travaux": travaux,
        "nb_cellules": nb_cell,
        "longueur_liaison_km": round(long_liaison, 2),
        "cout_total_keur": int(round(cout, -2)),
        "contribution_client_keur": int(round(
            cout * config["taux_contribution_client"], -2)),
        "delai_mois": config["delai_raccordement_mois"][domaine],
    })
    return solution


# ---------------------------------------------------------------------------
# Carte
# ---------------------------------------------------------------------------

def generer_carte(site, ouvrages, chemin_png):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        print("  [!] matplotlib absent : carte non générée")
        return None
    fig, ax = plt.subplots(figsize=(8, 7))
    couleurs = {400: "#d62728", 225: "#1f77b4", 150: "#17becf",
                90: "#2ca02c", 63: "#98df8a"}
    tensions_tracees = set()
    for ligne in ouvrages["lignes"]:
        t = ligne["tension_kv"]
        c = couleurs.get(t, "#999999")
        lbl = f"{t} kV" if t and t not in tensions_tracees else None
        tensions_tracees.add(t)
        lats = [p[0] for p in ligne["points"]]
        lons = [p[1] for p in ligne["points"]]
        ax.plot(lons, lats, color=c, lw=1.2, label=lbl, alpha=0.8)
    for poste in ouvrages["postes"][:30]:
        p = poste["points"][0]
        ax.plot(p[1], p[0], "s", color="#444444", ms=4)
    ax.plot(site["lon"], site["lat"], "*", color="#ff7f0e", ms=22,
            markeredgecolor="black", label="Site datacenter")
    ax.set_title(f"Réseau RTE autour de {site['commune']} "
                 f"({site.get('departement', '')})")
    ax.set_xlabel("Longitude")
    ax.set_ylabel("Latitude")
    ax.legend(loc="upper right", fontsize=8)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(chemin_png, dpi=150)
    plt.close(fig)
    return chemin_png


# ---------------------------------------------------------------------------
# Génération du document Word
# ---------------------------------------------------------------------------

def generer_docx(site, solution, ouvrages, config, chemin_carte, chemin_docx):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    def titre_page_garde(txt, size=14, bold=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.bold = bold
        r.font.size = Pt(size)
        return p

    def tableau_cle_valeur(donnees, largeur_cle=7):
        t = doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for cle, val in donnees:
            row = t.add_row()
            row.cells[0].text = cle
            row.cells[1].text = str(val)
            for run in row.cells[0].paragraphs[0].runs or \
                    [row.cells[0].paragraphs[0].add_run("")]:
                run.bold = True
            row.cells[0].width = Cm(largeur_cle)
        return t

    num_etude = site.get("numero_etude", "[À COMPLÉTER]")
    p_min = site.get("puissance_min_mw")
    p_max = site["puissance_max_mw"]
    puissance_txt = (f"{p_min} MW - {p_max} MW" if p_min and p_min != p_max
                     else f"{p_max} MW")

    # ---- Page de garde
    for _ in range(4):
        doc.add_paragraph()
    titre_page_garde(f"PRÉ-ÉTUDE EXPLORATOIRE n°{num_etude}", 16)
    titre_page_garde(f"DEMANDE DU CLIENT {site['client'].upper()}", 13)
    titre_page_garde("POUR LE RACCORDEMENT", 13)
    titre_page_garde(f"D'UN DATACENTER D'UNE PUISSANCE DE {puissance_txt}", 13)
    titre_page_garde(f"SUR LA COMMUNE DE {site['commune'].upper()} "
                     f"({site.get('departement', '')})", 13)
    doc.add_paragraph()
    avert = doc.add_paragraph()
    avert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = avert.add_run("Document interne indicatif établi au format des études "
                      "exploratoires RTE.\nCe document n'est pas émis par RTE "
                      "et ne l'engage en aucune manière.")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    r.italic = True
    doc.add_paragraph()
    doc.add_paragraph("Contacts :")
    tableau_cle_valeur([
        ("Destinataire Client", f"{site.get('contact_client', '[À COMPLÉTER]')}"),
        ("Rédacteur", f"{site.get('redacteur', '[À COMPLÉTER]')}"),
        ("Date", date.today().strftime("%d/%m/%Y")),
    ])
    doc.add_page_break()

    # ---- 1. Préambule
    doc.add_heading("1. Préambule", level=1)
    doc.add_heading("1.1 Exposé de la demande", level=2)
    tableau_cle_valeur([
        ("Nom du demandeur du raccordement", site["client"]),
        ("Date de la demande", site.get("date_demande", "[À COMPLÉTER]")),
        ("Localisation du Point de Raccordement (Commune et Département)",
         f"{site['commune']} ({site.get('departement', '')})"),
        ("Nature de la demande", "Nouveau raccordement"),
        ("Puissance de raccordement au soutirage", puissance_txt),
        ("Date de mise en service souhaitée",
         site.get("date_mes_souhaitee", "[À COMPLÉTER]")),
    ])

    doc.add_heading("1.2 Objectifs de l'étude", level=2)
    doc.add_paragraph(
        "La présente pré-étude exploratoire est établie selon la structure "
        "des études exploratoires réalisées par RTE conformément aux "
        "dispositions de l'article 1.2.2 de la documentation technique de "
        "référence de RTE.")
    doc.add_paragraph(
        "Cette étude simplifiée a pour objectif d'évaluer la faisabilité, le "
        "coût et le délai de mise en service du raccordement de l'installation "
        "du Client. Elle se limite à l'examen de la topologie du réseau public "
        "de transport au voisinage du site (elle n'examine pas les contraintes "
        "de transit, de court-circuit ou les perturbations électriques, qui "
        "relèvent des études menées par RTE).")
    doc.add_paragraph(
        "Cette étude a un caractère strictement indicatif. Seules l'étude "
        "exploratoire officielle demandée à RTE puis la Proposition Technique "
        "et Financière (PTF) constitueront des références opposables.")

    doc.add_heading("1.3 Hypothèses de l'étude", level=2)
    for h in [
        "Le réseau RTE pris en compte est le réseau existant tel que publié "
        "dans l'open data ODRE (lignes aériennes, lignes souterraines et "
        "postes électriques RTE), sans les évolutions décidées non publiées ;",
        f"Les éléments de coûts sont définis selon les conditions économiques "
        f"de l'année {config['conditions_economiques_annee']}, sur la base de "
        f"ratios paramétrables calibrés sur des études exploratoires RTE "
        f"récentes ;",
        "Le point de raccordement pris en compte est le centroïde du site "
        "communiqué par le Client ;",
        "L'arrivée d'autres installations de production et de consommation "
        "dans la zone d'étude serait de nature à remettre en cause les "
        "résultats de cette étude.",
    ]:
        doc.add_paragraph(h, style="List Bullet")

    doc.add_page_break()

    # ---- 2. Solutions de raccordement
    doc.add_heading("2. Solutions de raccordement HTB", level=1)
    doc.add_paragraph("Schéma de la zone d'étude et localisation de "
                      "l'installation :")
    if chemin_carte and os.path.exists(chemin_carte):
        doc.add_picture(chemin_carte, width=Cm(16))

    doc.add_heading("Domaine de tension de référence", level=2)
    doc.add_paragraph(
        "Comme défini dans l'article 105 de l'arrêté du 9 juin 2020 relatif "
        "aux prescriptions techniques de conception et de fonctionnement pour "
        "le raccordement aux réseaux d'électricité, le domaine de tension de "
        "raccordement de référence d'une installation de consommation est "
        "déterminé conformément au tableau ci-dessous :")
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Domaine de tension de raccordement de référence"
    t.rows[0].cells[1].text = ("Puissance de raccordement de l'installation "
                               "Prac inférieure à (MW)")
    for dom, seuil, libelle in DOMAINES_TENSION[:-1]:
        row = t.add_row()
        row.cells[0].text = f"{dom} ({libelle})"
        row.cells[1].text = str(seuil)
    row = t.add_row()
    row.cells[0].text = "HTB3 (400 kV)"
    row.cells[1].text = "au-delà de 400"
    doc.add_paragraph()
    doc.add_paragraph(
        f"Pour la puissance de raccordement demandée de {puissance_txt}, la "
        f"tension de raccordement de référence est la {solution['domaine']} "
        f"({solution['tension']}).")

    doc.add_heading("Contribution financière du Client", level=2)
    doc.add_paragraph(
        "Les principes d'établissement de la contribution du Client sont les "
        "suivants : le Client prend en charge 70 % des coûts des études et des "
        "travaux d'extension strictement nécessaires au raccordement, et 100 % "
        "des coûts des études et travaux supplémentaires demandés par le "
        "Client ; RTE prend en charge le reste (30 % extension, 100 % "
        "renforcement).")

    doc.add_heading(f"2.1 Solution : {solution['strategie']}", level=2)
    doc.add_paragraph("Description de la solution de raccordement :")
    tableau_cle_valeur([
        ("Tension de raccordement", solution["tension"]),
        ("Stratégie de raccordement", solution["strategie"]),
        ("Impact du raccordement sur le RPT",
         "[À COMPLÉTER après étude de transit RTE]"),
        ("Consistance des travaux", "\n".join(
            f"- {tr}" for tr in solution["travaux"])),
        ("Travaux de renforcement des ouvrages non dédiés du RPT",
         "[À COMPLÉTER après étude de transit RTE]"),
        ("Installation soumise à des limitations temporaires",
         "[À COMPLÉTER après étude de transit RTE]"),
        ("Coût du raccordement (indicatif, hors taxes)",
         f"Coût du raccordement : {solution['cout_total_keur']:,} k€\n"
         f"Contribution du client (estimée) : "
         f"{solution['contribution_client_keur']:,} k€".replace(",", " ")),
        ("Délai de raccordement (indicatif)",
         f"{solution['delai_mois']} mois à compter de la date d'acceptation "
         f"de la PTF"),
    ])
    doc.add_paragraph()
    doc.add_paragraph(
        "Nota : l'emplacement envisagé pour les ouvrages et le tracé potentiel "
        "du raccordement seront affinés par RTE au stade de l'étude "
        "exploratoire officielle puis de la PTF. Les coûts indiqués résultent "
        "de ratios paramétrables et ne constituent pas un engagement.")

    doc.add_page_break()

    # ---- 3. Notas
    doc.add_heading("3. Nota et points d'attention", level=1)
    doc.add_heading("3.1 Nota en lien avec le coût et la consistance "
                    "technique", level=2)
    for txt in [
        "L'étude ne prend pas en compte les éventuelles contraintes de tension "
        "et de court-circuit, y compris les éventuels moyens de compensation "
        "de réactif qui pourraient être nécessaires au raccordement de "
        "l'installation du Client.",
        "Le raccordement pourrait nécessiter la création de voies de "
        "transmission (fibre optique installée entre les postes du réseau). "
        "Ce besoin n'a pas été évalué à ce stade.",
    ]:
        doc.add_paragraph(txt, style="List Bullet")
    doc.add_heading("3.2 Nota en lien avec les démarches administratives",
                    level=2)
    for txt in [
        "Saisine de la Commission Nationale du Débat Public (CNDP) : en "
        "application de l'article L. 122-2 du code de l'environnement, "
        "l'ouvrage de raccordement constitue avec l'installation un même "
        "projet et doit être appréhendé dans sa globalité. Si le projet est "
        "soumis à une saisine de la CNDP, RTE devra co-saisir la CNDP en tant "
        "que co-maître d'ouvrage. Le délai indicatif ne prend pas en compte "
        "cette procédure.",
        "Évaluation environnementale : le projet de raccordement est soumis à "
        "une procédure dite de cas par cas (article R. 122-3 du code de "
        "l'environnement). Si l'une des composantes du projet est soumise à "
        "évaluation environnementale, le projet est, dans sa globalité, soumis "
        "à une évaluation environnementale commune et à une enquête publique "
        "commune.",
        f"Toute demande de PTF à RTE est accompagnée du versement d'une somme "
        f"forfaitaire de "
        f"{config['montant_forfaitaire_ptf_eur']:,} € HT".replace(",", " ")
        + ", déductible de la première échéance de la PTF.",
    ]:
        doc.add_paragraph(txt, style="List Bullet")

    # ---- 4. Conclusion
    doc.add_heading("4. Conclusion", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Solution de raccordement"
    hdr[1].text = "Coût à la charge du Client (indicatif)"
    hdr[2].text = "Délai de raccordement après acceptation PTF (indicatif)"
    row = t.add_row().cells
    row[0].text = solution["strategie"]
    row[1].text = (f"{solution['contribution_client_keur']:,} k€"
                   .replace(",", " "))
    row[2].text = f"{solution['delai_mois']} mois"
    doc.add_paragraph()
    doc.add_paragraph(
        "La présente pré-étude est une analyse interne indicative destinée à "
        "préparer la demande d'étude exploratoire officielle auprès de RTE. "
        "Si le Client souhaite poursuivre l'instruction de son raccordement, "
        "il doit déposer une demande d'étude exploratoire puis une demande de "
        "PTF auprès de RTE.")

    # ---- Annexe : ouvrages proches
    doc.add_page_break()
    doc.add_heading("ANNEXE : ouvrages RTE identifiés dans la zone d'étude",
                    level=1)
    doc.add_paragraph(f"Rayon de recherche : "
                      f"{config['rayon_recherche_km']} km. Source : open data "
                      f"ODRE (RTE). Distances à vol d'oiseau depuis le site.")
    doc.add_heading("Lignes HTB les plus proches", level=2)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for i, h in enumerate(["Ouvrage", "Tension", "Type", "Distance (km)"]):
        t.rows[0].cells[i].text = h
    deja = set()
    n = 0
    for ligne in ouvrages["lignes"]:
        cle = (ligne["nom"], ligne["tension_kv"])
        if cle in deja:
            continue
        deja.add(cle)
        row = t.add_row().cells
        row[0].text = ligne["nom"]
        row[1].text = f"{ligne['tension_kv']} kV" if ligne["tension_kv"] else "?"
        row[2].text = ligne["type"]
        row[3].text = f"{ligne['distance_km']:.1f}"
        n += 1
        if n >= 15:
            break
    doc.add_heading("Postes électriques les plus proches", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, h in enumerate(["Poste", "Tension max", "Distance (km)"]):
        t.rows[0].cells[i].text = h
    for poste in ouvrages["postes"][:15]:
        row = t.add_row().cells
        row[0].text = poste["nom"]
        row[1].text = (f"{poste['tension_kv']} kV" if poste["tension_kv"]
                       else "?")
        row[2].text = f"{poste['distance_km']:.1f}"

    doc.save(chemin_docx)
    return chemin_docx


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="DCEX - pré-étude exploratoire "
                                             "raccordement RTE datacenter")
    ap.add_argument("--site", required=True, help="Fichier JSON du site")
    ap.add_argument("--config", default=None, help="Fichier JSON de config")
    ap.add_argument("--offline", default=None,
                    help="Cache JSON réseau (mode hors-ligne, pour test)")
    ap.add_argument("--sortie", default=".", help="Dossier de sortie")
    args = ap.parse_args()

    ici = os.path.dirname(os.path.abspath(__file__))
    chemin_config = args.config or os.path.join(ici, "2026_07_13_dcex_config.json")
    with open(chemin_config, encoding="utf-8") as f:
        config = json.load(f)
    with open(args.site, encoding="utf-8") as f:
        site = json.load(f)

    # 1. Géolocalisation
    if "lat" in site and "lon" in site:
        print(f"Coordonnées fournies : {site['lat']}, {site['lon']}")
    else:
        print(f"Géocodage de {site['commune']} ...")
        c = geocoder_commune(site["commune"], site.get("departement"))
        site.update({"lat": c["lat"], "lon": c["lon"],
                     "departement": c["departement"], "commune": c["nom"]})
        print(f"  -> {c['nom']} ({c['departement']}) : "
              f"{c['lat']:.4f}, {c['lon']:.4f}")

    # 2. Collecte réseau
    print("Collecte des ouvrages RTE ...")
    reseau = collecter_reseau(site["lat"], site["lon"],
                              config["rayon_recherche_km"], args.offline)

    # 3. Analyse
    ouvrages = analyser_ouvrages(reseau, site["lat"], site["lon"])
    print(f"  {len(ouvrages['lignes'])} lignes, "
          f"{len(ouvrages['postes'])} postes dans le rayon")
    solution = construire_solution(site, ouvrages, config)
    print(f"  Tension de référence : {solution['domaine']} "
          f"({solution['tension']})")
    print(f"  Solution : {solution['strategie']}")
    print(f"  Coût indicatif : {solution['cout_total_keur']} k€, "
          f"délai {solution['delai_mois']} mois")

    # 4. Sorties
    os.makedirs(args.sortie, exist_ok=True)
    slug = strip_accents(site["commune"]).replace(" ", "_")
    prefixe = date.today().strftime("%Y_%m_%d")
    base = os.path.join(args.sortie,
                        f"{prefixe}_EE_{slug}_{site['puissance_max_mw']}MW")

    carte = generer_carte(site, ouvrages, base + "_carte.png")
    docx_path = generer_docx(site, solution, ouvrages, config, carte,
                             base + ".docx")
    donnees = {"site": {k: v for k, v in site.items()},
               "solution": {k: v for k, v in solution.items()
                            if k not in ("ligne_proche", "poste_proche")},
               "lignes_proches": [
                   {k: v for k, v in l.items() if k != "points"}
                   for l in ouvrages["lignes"][:20]],
               "postes_proches": [
                   {k: v for k, v in p.items() if k != "points"}
                   for p in ouvrages["postes"][:20]]}
    with open(base + "_donnees.json", "w", encoding="utf-8") as f:
        json.dump(donnees, f, ensure_ascii=False, indent=2)

    print(f"\nÉtude générée : {docx_path}")
    if carte:
        print(f"Carte : {carte}")
    print(f"Données : {base}_donnees.json")


if __name__ == "__main__":
    main()
